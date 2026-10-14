import html
import docx
from backend.importing.text_cleaning import sanitize_text
from backend.importing.structure import is_chapter_heading, is_scene_divider

def parse_docx(fpath: str) -> list[tuple[str, bool, str, str, str]]:
    doc = docx.Document(fpath)
    output = []
    for p in doc.paragraphs:
        txt = sanitize_text(p.text)
        if not txt:
            continue
        is_h = bool(p.style and p.style.name.lower().startswith(('heading', 'title', 'subtitle')))
        is_bold = False
        if p.runs and len(p.runs) > 0:
            is_bold = all(run.bold for run in p.runs if run.text.strip())

        is_ch, ch_title = is_chapter_heading(txt, is_heading_style=is_h, is_bold=is_bold)
        if is_ch:
            output.append((f"Chapter: {ch_title}", True, ch_title, "", ""))
        elif is_scene_divider(txt):
            output.append(("* * *", False, "", "", ""))
        else:
            formatted_runs = []
            for run in p.runs:
                if not run.text:
                    continue
                value = html.escape(run.text, quote=False)
                if run.italic:
                    value = f"<em>{value}</em>"
                if run.bold:
                    value = f"<strong>{value}</strong>"
                formatted_runs.append(value)
            output.append(("".join(formatted_runs) or html.escape(txt, quote=False), False, "", "", ""))
    return output

def inspect_docx_import(fpath: str) -> list[str]:
    document = docx.Document(fpath)
    warnings = []
    if document.tables:
        warnings.append(f"{len(document.tables)} table(s) found; table layout is not imported into the writing editor.")
    if document.inline_shapes:
        warnings.append(f"{len(document.inline_shapes)} image(s) found; images are not imported into the writing editor.")
    if document.part.element.xpath(".//w:hyperlink"):
        warnings.append("Hyperlinks were found; their text and link targets are not imported into the writing editor.")
    part_names = {part.partname.lower() for part in document.part.package.parts}
    if any("footnotes" in name or "endnotes" in name for name in part_names):
        warnings.append("Footnotes or endnotes were found; they are not imported into the writing editor.")
    if any("comments" in name for name in part_names):
        warnings.append("Word comments were found; they are not imported into the writing editor.")
    return warnings


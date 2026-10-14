import os
import re
import sys
import time
import uuid
import json
import html
import docx
import zipfile
import subprocess
import xml.etree.ElementTree as ET
from backend.database import get_db
from backend.text_utils import html_to_text

def sanitize_text(text: str) -> str:
    if not text:
        return ""
    text = text.replace('\x0b', '\n').replace('\r\n', '\n').replace('\r', '\n')
    text = re.sub(r'[\x00-\x08\x0e-\x1f]', '', text)
    return text.strip()

def strip_rtf_native(rtf_text: str, preserve_formatting: bool = False) -> str:
    """
    Cross-Platform RTF Stripper for Windows and Linux:
    Removes RTF formatting while decoding common ANSI and Unicode escapes.
    """
    if not rtf_text:
        return ""

    codepage_match = re.search(r'\\ansicpg(\d+)', rtf_text)
    encoding = f"cp{codepage_match.group(1)}" if codepage_match else "cp1252"
    output = []
    group_stack = []
    skip_group = False
    bold = False
    italic = False
    unicode_fallback_count = 1
    fallback_remaining = 0
    pending_high_surrogate = None
    index = 0

    def emit(text):
        if text:
            output.append((text, bold, italic))

    def append_unicode_unit(unit):
        nonlocal pending_high_surrogate
        if 0xD800 <= unit <= 0xDBFF:
            if pending_high_surrogate is not None:
                emit("\ufffd")
            pending_high_surrogate = unit
        elif 0xDC00 <= unit <= 0xDFFF and pending_high_surrogate is not None:
            emit(chr(0x10000 + ((pending_high_surrogate - 0xD800) << 10) + unit - 0xDC00))
            pending_high_surrogate = None
        else:
            if pending_high_surrogate is not None:
                emit("\ufffd")
                pending_high_surrogate = None
            emit(chr(unit))

    while index < len(rtf_text):
        char = rtf_text[index]
        if char == "{":
            prefix = rtf_text[index + 1:index + 80]
            destination = re.match(
                r"\\(?:\*\\)?([a-zA-Z]+)",
                prefix
            )
            destination_name = destination.group(1).lower() if destination else ""
            known_destinations = {
                "fonttbl", "colortbl", "stylesheet", "info", "pict", "object",
                "header", "headerl", "headerr", "footer", "footerl", "footerr",
                "footnote", "annotation", "datastore", "themedata",
                "colorschememapping",
            }
            starred_match = re.match(r"\\\*\\([a-zA-Z]+)", prefix)
            is_unknown_starred_destination = bool(starred_match and starred_match.group(1).lower() not in known_destinations)
            group_stack.append((skip_group, unicode_fallback_count, bold, italic))
            skip_group = skip_group or destination_name in known_destinations or is_unknown_starred_destination
            index += 1
            continue
        if char == "}":
            if group_stack:
                skip_group, unicode_fallback_count, bold, italic = group_stack.pop()
            index += 1
            continue
        if char != "\\":
            if fallback_remaining:
                fallback_remaining -= 1
            elif not skip_group:
                emit(char)
            index += 1
            continue

        index += 1
        if index >= len(rtf_text):
            break
        symbol = rtf_text[index]
        if symbol in "\\{}":
            if fallback_remaining:
                fallback_remaining -= 1
            elif not skip_group:
                emit(symbol)
            index += 1
            continue
        if symbol == "'":
            match = re.match(r"'([0-9a-fA-F]{2})", rtf_text[index:])
            if match:
                if fallback_remaining:
                    fallback_remaining -= 1
                elif not skip_group:
                    try:
                        emit(bytes.fromhex(match.group(1)).decode(encoding, errors="replace"))
                    except LookupError:
                        emit(bytes.fromhex(match.group(1)).decode("cp1252", errors="replace"))
                index += len(match.group(0))
            else:
                index += 1
            continue
        if not symbol.isalpha():
            if not skip_group and not fallback_remaining:
                emit({"~": "\u00a0", "_": "\u2011", "-": ""}.get(symbol, ""))
            elif fallback_remaining:
                fallback_remaining -= 1
            index += 1
            continue

        start = index
        while index < len(rtf_text) and rtf_text[index].isalpha():
            index += 1
        word = rtf_text[start:index].lower()
        number_match = re.match(r"-?\d+", rtf_text[index:])
        number = int(number_match.group(0)) if number_match else None
        if number_match:
            index += len(number_match.group(0))
        if index < len(rtf_text) and rtf_text[index] == " ":
            index += 1

        if word == "uc" and number is not None:
            unicode_fallback_count = max(0, number)
        elif word == "u" and number is not None:
            if not skip_group and not fallback_remaining:
                append_unicode_unit(number if number >= 0 else number + 65536)
            fallback_remaining = unicode_fallback_count
        elif word == "b":
            bold = number != 0
        elif word == "i":
            italic = number != 0
        elif word in ("par", "line") and not skip_group:
            emit("\n")
        elif word == "tab" and not skip_group:
            emit("\t")
        elif word in ("emdash", "endash", "lquote", "rquote", "ldblquote", "rdblquote") and not skip_group:
            emit({
                "emdash": "\u2014", "endash": "\u2013", "lquote": "\u2018",
                "rquote": "\u2019", "ldblquote": "\u201c", "rdblquote": "\u201d",
            }[word])
        elif word == "bin" and number is not None:
            index += max(0, number)

    if pending_high_surrogate is not None:
        emit("\ufffd")

    if not preserve_formatting:
        return sanitize_text("".join(text for text, _, _ in output))

    paragraphs = [[]]
    for text, token_bold, token_italic in output:
        fragments = text.split("\n")
        for fragment_index, fragment in enumerate(fragments):
            if fragment:
                paragraphs[-1].append((fragment, token_bold, token_italic))
            if fragment_index < len(fragments) - 1:
                paragraphs.append([])

    rendered_paragraphs = []
    for paragraph in paragraphs:
        runs = []
        for text, token_bold, token_italic in paragraph:
            if runs and runs[-1][1:] == (token_bold, token_italic):
                runs[-1] = (runs[-1][0] + text, token_bold, token_italic)
            else:
                runs.append((text, token_bold, token_italic))
        rendered = []
        for text, token_bold, token_italic in runs:
            escaped = html.escape(text, quote=False)
            if token_bold and token_italic:
                escaped = f"<strong><em>{escaped}</em></strong>"
            elif token_bold:
                escaped = f"<strong>{escaped}</strong>"
            elif token_italic:
                escaped = f"<em>{escaped}</em>"
            rendered.append(escaped)
        rendered_paragraphs.append("".join(rendered))
    return "\n\n".join(rendered_paragraphs).strip()

def read_file_content(fp: str, preserve_formatting: bool = False) -> str:
    """Reads RTF, TXT, or MD files across macOS, Windows 10, and Linux."""
    if not fp or not os.path.exists(fp):
        return ""
    if fp.lower().endswith('.rtf'):
        if preserve_formatting:
            with open(fp, 'r', encoding='utf-8', errors='replace') as f:
                return strip_rtf_native(f.read(), preserve_formatting=True)
        if sys.platform == 'darwin':
            try:
                res = subprocess.run(['textutil', '-convert', 'txt', fp, '-stdout'], capture_output=True, text=True)
                if res.returncode == 0 and res.stdout:
                    return res.stdout
            except OSError:
                pass
        with open(fp, 'r', encoding='utf-8', errors='replace') as f:
            return strip_rtf_native(f.read())

    if fp.lower().endswith(('.txt', '.md')):
        with open(fp, 'r', encoding='utf-8', errors='replace') as f:
            return f.read()
    return ""

def normalize_paragraphs(raw_text: str) -> list[str]:
    cleaned = sanitize_text(raw_text)
    if not cleaned:
        return []
    if "\n\n" in cleaned:
        chunks = re.split(r'\n{2,}', cleaned)
    else:
        chunks = cleaned.split("\n")
    paragraphs = []
    for c in chunks:
        p = re.sub(r'^[ \t]+', '', c).strip()
        if p and not p.startswith("{\\rtf"):
            paragraphs.append(p)
    return paragraphs

def _safe_inline_markup(value: str) -> str:
    pieces = re.split(r"(<\/?(?:strong|em)>)", value, flags=re.IGNORECASE)
    return "".join(
        piece if re.fullmatch(r"<\/?(?:strong|em)>", piece, flags=re.IGNORECASE)
        else html.escape(html.unescape(piece), quote=False)
        for piece in pieces
    )

def is_chapter_heading(text: str, is_heading_style: bool = False, is_bold: bool = False) -> tuple[bool, str]:
    t = text.strip()
    if not t or len(t) > 85:
        return False, ""
    
    clean = re.sub(r"[\*#_~]", "", t).strip()
    clean_lower = clean.lower()

    if clean.isdigit() and int(clean) < 150:
        return True, f"Chapter {clean}"
    if re.match(r"^(?:[IVXLCDM]+)\.?$", clean, re.IGNORECASE) and len(clean) <= 6:
        return True, f"Chapter {clean.upper()}"

    words = ["one","two","three","four","five","six","seven","eight","nine","ten","eleven","twelve","thirteen","fourteen","fifteen","sixteen","seventeen","eighteen","nineteen","twenty"]
    if clean_lower in words:
        return True, f"Chapter {clean.title()}"

    match = re.match(r"^(?:chapter|act|part|prologue|epilogue|interlude|book)\b(?:\s+[\w\d]+)?(?:\s*[:\-–—]\s*(.*))?$", clean, re.I)
    if match:
        return True, clean

    if (is_heading_style or is_bold) and len(clean) < 55:
        if not clean.endswith(('.', '!', '?')) and len(clean.split()) <= 7:
            return True, clean

    return False, ""

def is_scene_divider(text: str) -> bool:
    t = text.strip()
    return bool(re.match(r"^(?:(?:\*\s*){3,}|(?:#\s*){3,}|(?:~\s*){3,}|(?:•\s*){3,}|-{3,})$", t))

def derive_smart_scene_title(first_para: str, sc_num: int, ch_title: str, custom_title: str | None = None) -> str:
    if custom_title and custom_title.strip() and not custom_title.lower().startswith(("scene ", "untitled")):
        return custom_title.strip()
    sub_match = re.search(r"[:\-–—]\s*(.+)$", ch_title)
    if sc_num == 1 and sub_match:
        return sub_match.group(1).strip()
    if first_para:
        first_sentence = re.split(r'[.!?]', first_para)[0].strip()
        words = first_sentence.split()
        if len(words) >= 2:
            candidate = " ".join(words[:5]).strip('.,;:"!?')
            candidate = re.sub(r'^(and|but|then|so|because|when|the|a|an)\s+', '', candidate, flags=re.I)
            if len(candidate) > 3:
                return candidate.title()
    return f"Scene {sc_num}"

def parse_docx(fpath: str) -> list[tuple[str, bool, str, str, str]]:
    doc = docx.Document(fpath)
    output = []
    for p in doc.paragraphs:
        txt = sanitize_text(p.text)
        if not txt:
            continue
        is_h = bool(p.style and p.style.name.lower().startswith(('heading', 'title', 'subtitle')))
        is_bold = False
        if p.runs and len(p.runs) > 0:
            is_bold = all(run.bold for run in p.runs if run.text.strip())

        is_ch, ch_title = is_chapter_heading(txt, is_heading_style=is_h, is_bold=is_bold)
        if is_ch:
            output.append((f"Chapter: {ch_title}", True, ch_title, "", ""))
        elif is_scene_divider(txt):
            output.append(("* * *", False, "", "", ""))
        else:
            formatted_runs = []
            for run in p.runs:
                if not run.text:
                    continue
                value = html.escape(run.text, quote=False)
                if run.italic:
                    value = f"<em>{value}</em>"
                if run.bold:
                    value = f"<strong>{value}</strong>"
                formatted_runs.append(value)
            output.append(("".join(formatted_runs) or html.escape(txt, quote=False), False, "", "", ""))
    return output

def inspect_docx_import(fpath: str) -> list[str]:
    document = docx.Document(fpath)
    warnings = []
    if document.tables:
        warnings.append(f"{len(document.tables)} table(s) found; table layout is not imported into the writing editor.")
    if document.inline_shapes:
        warnings.append(f"{len(document.inline_shapes)} image(s) found; images are not imported into the writing editor.")
    if document.part.element.xpath(".//w:hyperlink"):
        warnings.append("Hyperlinks were found; their text and link targets are not imported into the writing editor.")
    part_names = {part.partname.lower() for part in document.part.package.parts}
    if any("footnotes" in name or "endnotes" in name for name in part_names):
        warnings.append("Footnotes or endnotes were found; they are not imported into the writing editor.")
    if any("comments" in name for name in part_names):
        warnings.append("Word comments were found; they are not imported into the writing editor.")
    return warnings

def build_import_report(manifest: dict, warnings: list[str] | None = None) -> dict:
    tuples = manifest.get("manuscript", [])
    scene_titles = [item for item in tuples if item[0].startswith("### SCENE_TITLE: ")]
    outline_only = 0
    scene_count = 0
    scene_open = False
    synopses = 0
    notes = 0
    prose_count = 0
    for position, item in enumerate(tuples):
        text = item[0]
        is_scene_title = text.startswith("### SCENE_TITLE: ")
        is_chapter = text.startswith("Chapter: ")
        synopsis = item[3].strip() if len(item) > 3 and isinstance(item[3], str) else ""
        note = item[4].strip() if len(item) > 4 and isinstance(item[4], str) else ""
        if is_scene_title or is_chapter:
            synopses += bool(synopsis)
            notes += bool(note)
        if is_chapter:
            scene_open = False
            continue
        if is_scene_title:
            scene_count += 1
            scene_open = True
        elif is_scene_divider(text):
            scene_open = False
        elif text.strip():
            prose_count += 1
            if not scene_open:
                scene_count += 1
                scene_open = True

    scene_count = max(len(scene_titles), scene_count)
    for position, item in enumerate(tuples):
        if not item[0].startswith("### SCENE_TITLE: "):
            continue
        has_prose = False
        for entry in tuples[position + 1:]:
            if entry[0].startswith(("### SCENE_TITLE: ", "Chapter: ")):
                break
            if not is_scene_divider(entry[0]) and entry[0].strip():
                has_prose = True
                break
        if not has_prose:
            outline_only += 1
    paragraphs = [
        item[0] for item in tuples
        if not item[0].startswith(("### SCENE_TITLE: ", "Chapter: "))
        and not is_scene_divider(item[0])
    ]
    formatted_count = sum(
        1 for paragraph in paragraphs
        if re.search(r"</?(?:strong|em)>", paragraph, re.IGNORECASE)
    )
    word_count = len(html_to_text("\n".join(paragraphs)).split())
    report_warnings = list(warnings or [])
    report_warnings.extend(manifest.get("import_warnings", []))
    return {
        "chapters": sum(item[0].startswith("Chapter: ") for item in tuples),
        "scenes": scene_count,
        "outline_only": outline_only,
        "paragraphs": len(paragraphs),
        "words": word_count,
        "formatted_paragraphs": formatted_count,
        "characters": len(manifest.get("characters", [])),
        "locations": len(manifest.get("locations", [])),
        "lore": len(manifest.get("lore", [])),
        "synopses": synopses,
        "notes": notes,
        "warnings": list(dict.fromkeys(report_warnings)),
    }

def parse_plaintext(fpath: str) -> list[tuple[str, bool, str, str, str]]:
    with open(fpath, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read()
        output = []
        for p in normalize_paragraphs(content):
            is_ch, ch_title = is_chapter_heading(p)
            if is_ch:
                output.append((f"Chapter: {ch_title}", True, ch_title, "", ""))
            elif is_scene_divider(p):
                output.append(("* * *", False, "", "", ""))
            else:
                output.append((p, False, "", "", ""))
        return output

def parse_scrivener_project(scriv_path: str) -> dict:
    files_map = {}
    legacy_doc = re.compile(
        r"^(\d+)(?:_(synopsis|notes?|annotations?|comments?|references))?\.(?:rtf|txt)$", re.IGNORECASE
    )
    role_names = {
        "synopsis": "synopsis", "note": "notes", "notes": "notes",
        "annotation": "annotations", "annotations": "annotations",
        "comment": "comments", "comments": "comments", "references": "references",
    }
    for root, _, files in os.walk(scriv_path):
        relative_root = os.path.relpath(root, scriv_path).lower()
        if any(x in relative_root for x in ['settings', 'quicklook', 'snapshots']):
            continue
        data_id = os.path.basename(root)
        for f in files:
            full_p = os.path.join(root, f)
            fl = f.lower()
            legacy = legacy_doc.match(f) if data_id.lower() == "docs" else None
            if legacy:
                role = role_names.get((legacy.group(2) or "").lower(), "content")
                files_map.setdefault(legacy.group(1), {})[role] = full_p
            elif fl in ('content.rtf', 'content.txt'):
                files_map.setdefault(data_id, {})['content'] = full_p
            elif fl in ('synopsis.rtf', 'synopsis.txt'):
                files_map.setdefault(data_id, {})['synopsis'] = full_p
            elif fl in ('notes.rtf', 'notes.txt', 'note.rtf', 'note.txt'):
                files_map.setdefault(data_id, {})['notes'] = full_p
            elif fl in ('annotations.rtf', 'annotations.txt', 'annotation.rtf', 'annotation.txt'):
                files_map.setdefault(data_id, {})['annotations'] = full_p
            elif fl in ('comments.rtf', 'comments.txt', 'comment.rtf', 'comment.txt'):
                files_map.setdefault(data_id, {})['comments'] = full_p
            elif fl in ('references.rtf', 'references.txt'):
                files_map.setdefault(data_id, {})['references'] = full_p

    scrivx_files = sorted(f for f in os.listdir(scriv_path) if f.lower().endswith('.scrivx'))
    manuscript_tuples = []
    characters = []
    locations = []
    lore_entries = []
    project_metadata = {}

    if scrivx_files:
        scrivx_path = os.path.join(scriv_path, scrivx_files[0])
        try:
            tree = ET.parse(scrivx_path)
            root = tree.getroot()
            project_metadata.update({
                "scrivx_file": os.path.basename(scrivx_path),
                "project_attributes": dict(root.attrib),
            })

            def xml_value(element):
                children = list(element)
                value = (element.text or "").strip()
                if not children and not element.attrib:
                    return value
                result = {}
                if element.attrib:
                    result["@attributes"] = dict(element.attrib)
                if value:
                    result["#text"] = value
                for child in children:
                    child_value = xml_value(child)
                    if child.tag in result:
                        if not isinstance(result[child.tag], list):
                            result[child.tag] = [result[child.tag]]
                        result[child.tag].append(child_value)
                    else:
                        result[child.tag] = child_value
                return result

            def item_metadata(item_node, title, uuid_val, itype, source_path):
                metadata = {}
                for child in item_node:
                    if child.tag in ("Title", "Children"):
                        continue
                    value = xml_value(child)
                    if child.tag in metadata:
                        if not isinstance(metadata[child.tag], list):
                            metadata[child.tag] = [metadata[child.tag]]
                        metadata[child.tag].append(value)
                    else:
                        metadata[child.tag] = value
                return {
                    "uuid": uuid_val,
                    "type": itype,
                    "title": title,
                    "source_path": source_path,
                    "metadata": metadata,
                }

            project_metadata["metadata"] = {
                child.tag: xml_value(child)
                for child in root
                if child.tag != "Binder"
            }

            def attach_auxiliary_files(metadata, bundle):
                extras = {
                    role: read_file_content(bundle[role])
                    for role in ("annotations", "comments", "references")
                    if bundle.get(role)
                }
                if extras:
                    metadata["external_files"] = extras
                return metadata

            def traverse_folder(item_node, path=()):
                nonlocal manuscript_tuples, characters, locations, lore_entries
                title_node = item_node.find('Title')
                raw_title = title_node.text.strip() if title_node is not None and title_node.text else ""
                title_lower = raw_title.lower()
                itype = item_node.attrib.get('Type', '')
                uuid_val = str(
                    item_node.attrib.get('UUID') or item_node.attrib.get('ID')
                    or item_node.attrib.get('Id') or ""
                )
                item_path = (*path, raw_title) if raw_title else path
                source_metadata = item_metadata(item_node, raw_title, uuid_val, itype, list(item_path))

                is_char_folder = "character" in title_lower or "cast" in title_lower or "people" in title_lower
                is_loc_folder = "place" in title_lower or "location" in title_lower or "setting" in title_lower
                is_lore_folder = "research" in title_lower or "lore" in title_lower or "glossary" in title_lower or "world" in title_lower

                file_bundle = files_map.get(uuid_val, {})
                source_metadata = attach_auxiliary_files(source_metadata, file_bundle)
                content_text = read_file_content(file_bundle.get('content', ''), preserve_formatting=True)
                synopsis_text = read_file_content(file_bundle.get('synopsis', '')) or (item_node.findtext('Synopsis') or '').strip()
                notes_text = read_file_content(file_bundle.get('notes', '')) or (item_node.findtext('Notes') or '').strip()
                if itype.lower() == "folder":
                    project_metadata.setdefault("binder_folders", []).append(source_metadata)

                if is_char_folder:
                    def collect_characters(folder_node, folder_path):
                        children_node = folder_node.find('Children')
                        for child in children_node.findall('BinderItem') if children_node is not None else []:
                            child_title_node = child.find('Title')
                            child_title = child_title_node.text.strip() if child_title_node is not None and child_title_node.text else ""
                            child_uuid = str(child.attrib.get('UUID') or child.attrib.get('Id') or "")
                            child_type = child.attrib.get('Type', '')
                            child_path = (*folder_path, child_title) if child_title else folder_path
                            child_meta = item_metadata(child, child_title, child_uuid, child_type, list(child_path))
                            child_bundle = files_map.get(child_uuid, {})
                            child_meta = attach_auxiliary_files(child_meta, child_bundle)
                            child_content = read_file_content(child_bundle.get('content', ''))
                            child_notes = read_file_content(child_bundle.get('notes', '')) or (child.findtext('Notes') or '').strip()
                            child_synopsis = read_file_content(child_bundle.get('synopsis', '')) or (child.findtext('Synopsis') or '').strip()
                            grandchildren = child.find('Children')
                            if child_type == 'Folder' or grandchildren is not None:
                                collect_characters(child, child_path)
                            elif child_title:
                                characters.append({
                                    "name": child_title,
                                    "bio": child_content or child_notes,
                                    "synopsis": child_synopsis,
                                    "notes": child_notes,
                                    "metadata": child_meta,
                                })
                    collect_characters(item_node, item_path)
                    return

                if is_loc_folder:
                    def collect_locations(folder_node, folder_path):
                        children_node = folder_node.find('Children')
                        for child in children_node.findall('BinderItem') if children_node is not None else []:
                            child_title_node = child.find('Title')
                            child_title = child_title_node.text.strip() if child_title_node is not None and child_title_node.text else ""
                            child_uuid = str(child.attrib.get('UUID') or child.attrib.get('Id') or "")
                            child_type = child.attrib.get('Type', '')
                            child_path = (*folder_path, child_title) if child_title else folder_path
                            child_meta = item_metadata(child, child_title, child_uuid, child_type, list(child_path))
                            child_bundle = files_map.get(child_uuid, {})
                            child_meta = attach_auxiliary_files(child_meta, child_bundle)
                            child_content = read_file_content(child_bundle.get('content', ''))
                            child_notes = read_file_content(child_bundle.get('notes', '')) or (child.findtext('Notes') or '').strip()
                            child_synopsis = read_file_content(child_bundle.get('synopsis', '')) or (child.findtext('Synopsis') or '').strip()
                            grandchildren = child.find('Children')
                            if child_type == 'Folder' or grandchildren is not None:
                                collect_locations(child, child_path)
                            elif child_title:
                                locations.append({
                                    "name": child_title,
                                    "desc": child_content or child_notes,
                                    "synopsis": child_synopsis,
                                    "notes": child_notes,
                                    "metadata": child_meta,
                                })
                    collect_locations(item_node, item_path)
                    return

                if is_lore_folder:
                    def collect_lore(folder_node, folder_path):
                        children_node = folder_node.find('Children')
                        for child in children_node.findall('BinderItem') if children_node is not None else []:
                            child_title_node = child.find('Title')
                            child_title = child_title_node.text.strip() if child_title_node is not None and child_title_node.text else ""
                            child_uuid = str(child.attrib.get('UUID') or child.attrib.get('Id') or "")
                            child_type = child.attrib.get('Type', '')
                            child_path = (*folder_path, child_title) if child_title else folder_path
                            child_meta = item_metadata(child, child_title, child_uuid, child_type, list(child_path))
                            child_bundle = files_map.get(child_uuid, {})
                            child_meta = attach_auxiliary_files(child_meta, child_bundle)
                            child_content = read_file_content(child_bundle.get('content', ''))
                            child_notes = read_file_content(child_bundle.get('notes', '')) or (child.findtext('Notes') or '').strip()
                            child_synopsis = read_file_content(child_bundle.get('synopsis', '')) or (child.findtext('Synopsis') or '').strip()
                            grandchildren = child.find('Children')
                            if child_type == 'Folder' or grandchildren is not None:
                                collect_lore(child, child_path)
                            elif child_title:
                                definition = child_content or child_synopsis or child_notes
                                lore_entries.append({
                                    "term": child_title,
                                    "definition": definition,
                                    "notes": child_notes,
                                    "metadata": child_meta,
                                })
                    collect_lore(item_node, item_path)
                    return

                children = item_node.findall('Children/BinderItem') if item_node.find('Children') is not None else []
                if raw_title.lower() in ('draft', 'manuscript') and children:
                    for child in children:
                        traverse_folder(child, item_path)
                    return
                if itype == 'Folder' or (children and itype != 'Text'):
                    folder_name = raw_title or "Chapter"
                    manuscript_tuples.append((f"Chapter: {folder_name}", True, folder_name, synopsis_text, notes_text, source_metadata))
                    for child in children:
                        traverse_folder(child, item_path)
                elif itype in ['Text', 'Document'] or not children:
                    doc_title = raw_title or "Scene"
                    paras = normalize_paragraphs(content_text)
                    if paras or synopsis_text.strip() or notes_text.strip():
                        manuscript_tuples.append((f"### SCENE_TITLE: {doc_title}", False, doc_title, synopsis_text, notes_text, source_metadata))
                        for p in paras:
                            manuscript_tuples.append((p, False, doc_title, synopsis_text, notes_text))

            binder_node = root.find('Binder')
            target_root = binder_node if binder_node is not None else root
            top_items = target_root.findall('BinderItem')
            if not top_items:
                children_node = target_root.find('Children')
                top_items = children_node.findall('BinderItem') if children_node is not None else []
            for top_item in top_items:
                traverse_folder(top_item)

        except ET.ParseError as e:
            raise ValueError(f"Could not read Scrivener project metadata from {os.path.basename(scrivx_path)}: {e}") from e

    if not manuscript_tuples and files_map:
        for key, bundle in sorted(files_map.items()):
            content_text = read_file_content(bundle.get('content', ''), preserve_formatting=True)
            synopsis_text = read_file_content(bundle.get('synopsis', ''))
            notes_text = read_file_content(bundle.get('notes', ''))
            paras = normalize_paragraphs(content_text)
            if paras or synopsis_text.strip() or notes_text.strip():
                manuscript_tuples.append((f"### SCENE_TITLE: {key}", False, str(key), synopsis_text, notes_text))
                for p in paras:
                    manuscript_tuples.append((p, False, str(key), synopsis_text, notes_text))

    return {
        "manuscript": manuscript_tuples,
        "characters": characters,
        "locations": locations,
        "lore": lore_entries,
        "scrivener": project_metadata,
    }

def build_binder_from_manifest(manifest: dict, book_title: str) -> str:
    tuples = manifest.get("manuscript", [])
    characters = manifest.get("characters", [])
    locations = manifest.get("locations", [])
    lore_entries = manifest.get("lore", [])

    if not tuples:
        raise ValueError(f"No narrative prose could be extracted from '{book_title}'.")

    conn = get_db()
    c = conn.cursor()
    now = time.strftime("%Y-%m-%d %H:%M:%S")

    book_id = f"book_{uuid.uuid4().hex[:8]}"
    book_card = json.dumps({"scrivener": manifest.get("scrivener", {})}, ensure_ascii=False)
    c.execute("""
        INSERT INTO binder_nodes (id, parent_id, project_id, node_type, title, sort_order, synopsis, content, notes, status, word_count, is_archived, epigraph, card_data, created_at, updated_at)
        VALUES (?, NULL, 'default', 'BOOK', ?, 10.0, 'Imported Manuscript', '', '', 'DRAFT', 0, 0, '', ?, ?, ?)
    """, (book_id, book_title, book_card, now, now))

    current_ch_id = None
    current_sc_id = None
    current_ch_title = "Chapter 1"
    current_custom_title = None
    current_synopsis = ""
    current_notes = ""
    current_scene_metadata = {}
    current_chapter_metadata = {}
    ch_sort = 10.0
    sc_sort = 10.0
    sc_num = 0
    current_scene_paras = []
    first_created_scene_id = None

    def flush_scene():
        nonlocal current_sc_id, current_scene_paras, first_created_scene_id, current_custom_title, current_synopsis, current_notes, current_scene_metadata
        if current_sc_id:
            full_text = "\n".join(
                f"<p>{_safe_inline_markup(paragraph)}</p>"
                for paragraph in current_scene_paras
            ).strip()
            wc = len(html_to_text(full_text).split())
            plain_text = html_to_text(full_text)
            final_synopsis = current_synopsis.strip() if current_synopsis.strip() else ((plain_text[:140] + "...") if len(plain_text) > 140 else plain_text)
            first_para = html_to_text(current_scene_paras[0]) if current_scene_paras else ""
            final_title = derive_smart_scene_title(first_para, sc_num, current_ch_title, current_custom_title)
            imported_card = json.dumps({"scrivener": current_scene_metadata}, ensure_ascii=False) if current_scene_metadata else "{}"
            
            c.execute("""
                UPDATE binder_nodes 
                SET content = ?, word_count = ?, synopsis = ?, notes = ?, title = ?, card_data = ?
                WHERE id = ?
            """, (full_text, wc, final_synopsis, current_notes, final_title, imported_card, current_sc_id))
            
            if not first_created_scene_id:
                first_created_scene_id = current_sc_id
                
            current_scene_paras = []
            current_custom_title = None
            current_synopsis = ""
            current_notes = ""
            current_scene_metadata = {}
            current_sc_id = None

    def ensure_chapter():
        nonlocal current_ch_id
        if current_ch_id:
            return
        current_ch_id = f"ch_{uuid.uuid4().hex[:8]}"
        c.execute("""
            INSERT INTO binder_nodes (id, parent_id, project_id, node_type, title, sort_order, synopsis, content, notes, status, word_count, is_archived, epigraph, card_data, created_at, updated_at)
            VALUES (?, ?, 'default', 'CHAPTER', 'Chapter 1', 10.0, '', '', '', 'DRAFT', 0, 0, '', '{}', ?, ?)
        """, (current_ch_id, book_id, now, now))

    def ensure_scene():
        nonlocal current_sc_id, sc_num, sc_sort, current_scene_metadata
        if current_sc_id:
            return
        ensure_chapter()
        sc_num += 1
        current_sc_id = f"scene_{uuid.uuid4().hex[:8]}"
        placeholder_title = current_custom_title or f"Scene {sc_num}"
        c.execute("""
            INSERT INTO binder_nodes (id, parent_id, project_id, node_type, title, sort_order, synopsis, content, notes, status, word_count, is_archived, epigraph, card_data, created_at, updated_at)
            VALUES (?, ?, 'default', 'SCENE', ?, ?, '', '', '', 'DRAFT', 0, 0, '', ?, ?, ?)
        """, (current_sc_id, current_ch_id, placeholder_title, sc_sort, json.dumps({"scrivener": current_scene_metadata}, ensure_ascii=False) if current_scene_metadata else "{}", now, now))
        sc_sort += 10.0

    for item in tuples:
        text = item[0]
        is_h = item[1]
        custom_sc_title = item[2] if len(item) > 2 else None
        item_synopsis = item[3] if len(item) > 3 else ""
        item_notes = item[4] if len(item) > 4 else ""
        item_scrivener_metadata = item[5] if len(item) > 5 and isinstance(item[5], dict) else {}

        t_clean = text.strip()
        is_ch, ch_heading = is_chapter_heading(t_clean, is_heading_style=is_h)

        if t_clean.startswith("### SCENE_TITLE: "):
            flush_scene()
            current_custom_title = t_clean.replace("### SCENE_TITLE: ", "").strip()
            current_synopsis = item_synopsis
            current_notes = item_notes
            current_scene_metadata = item_scrivener_metadata
            ensure_scene()
            continue

        if is_ch or t_clean.startswith("Chapter: "):
            flush_scene()
            current_sc_id = None
            sc_num = 0
            sc_sort = 10.0
            current_ch_id = f"ch_{uuid.uuid4().hex[:8]}"
            current_ch_title = ch_heading if is_ch else t_clean.replace("Chapter: ", "").strip()
            current_chapter_metadata = item_scrivener_metadata
            
            c.execute("""
                INSERT INTO binder_nodes (id, parent_id, project_id, node_type, title, sort_order, synopsis, content, notes, status, word_count, is_archived, epigraph, card_data, created_at, updated_at)
                VALUES (?, ?, 'default', 'CHAPTER', ?, ?, ?, '', ?, 'DRAFT', 0, 0, '', ?, ?, ?)
            """, (current_ch_id, book_id, current_ch_title, ch_sort, item_synopsis, item_notes, json.dumps({"scrivener": current_chapter_metadata}, ensure_ascii=False), now, now))
            ch_sort += 10.0
            continue

        if is_scene_divider(t_clean):
            flush_scene()
            continue

        current_scene_metadata = item_scrivener_metadata or current_scene_metadata
        current_synopsis = item_synopsis or current_synopsis
        current_notes = item_notes or current_notes
        if not current_custom_title and custom_sc_title:
            current_custom_title = custom_sc_title
        ensure_scene()

        current_scene_paras.append(t_clean)

    flush_scene()

    # Commit Character Sheets
    for ch in characters:
        name_clean = ch["name"].strip()
        if name_clean and len(name_clean) >= 2:
            eid = re.sub(r'\W+', '_', name_clean.lower())
            axioms = json.dumps({
                "bio": ch.get("bio", ""),
                "synopsis": ch.get("synopsis", ""),
                "notes": ch.get("notes", ""),
                "scrivener": ch.get("metadata", {}),
            }, ensure_ascii=False)
            sensory = ch.get("synopsis", "")[:200]
            c.execute("""
                INSERT INTO canonical_entities (entity_id, name, entity_type, role, aliases, sensory_profile, status, axioms, created_at)
                VALUES (?, ?, 'CHARACTER', 'SUPPORTING', '', ?, 'ALIVE', ?, ?)
                ON CONFLICT(entity_id) DO UPDATE SET 
                    axioms = excluded.axioms,
                    sensory_profile = CASE WHEN canonical_entities.sensory_profile = '' THEN excluded.sensory_profile ELSE canonical_entities.sensory_profile END
            """, (eid, name_clean, sensory, axioms, now))

    # Commit Locations
    for loc in locations:
        loc_name = loc["name"].strip()
        if loc_name:
            eid = f"loc_{re.sub(r'\W+', '_', loc_name.lower())}"
            axioms = json.dumps({
                "description": loc.get("desc", ""),
                "synopsis": loc.get("synopsis", ""),
                "notes": loc.get("notes", ""),
                "scrivener": loc.get("metadata", {}),
            }, ensure_ascii=False)
            c.execute("""
                INSERT INTO canonical_entities (entity_id, name, entity_type, role, aliases, sensory_profile, status, axioms, created_at)
                VALUES (?, ?, 'LOCATION', 'SUPPORTING', '', '', 'ALIVE', ?, ?)
                ON CONFLICT(entity_id) DO NOTHING
            """, (eid, loc_name, axioms, now))

    # Commit Lore
    for rule in lore_entries:
        term = rule["term"].strip()
        definition = rule["definition"].strip()
        if term and definition:
            c.execute("""
                INSERT OR REPLACE INTO style_sheet (term, category, rule_definition, pronunciation, metadata_json, created_at)
                VALUES (?, 'LORE', ?, '', ?, ?)
            """, (term, definition, json.dumps({
                "notes": rule.get("notes", ""),
                "scrivener": rule.get("metadata", {}),
            }, ensure_ascii=False), now))

    conn.commit()
    conn.close()
    return first_created_scene_id or book_id

def scan_local_manuscripts(with_status: bool = False):
    found = []
    folders = []
    ignored = {'venv', '.venv', 'site-packages', 'temp', 'node_modules', '.git'}
    home = os.path.expanduser("~")
    for name in ("Documents", "Desktop", "Downloads"):
        folder = os.path.join(home, name)
        status = {"folder": folder, "exists": os.path.isdir(folder), "readable": False}
        if status["exists"]:
            try:
                os.listdir(folder)
                status["readable"] = True
            except OSError:
                pass
        folders.append(status)
        if status["readable"]:
            found.extend(_scan_manuscript_folder(folder, home, ignored))
    found.sort(key=lambda x: x["title"].lower())
    found = found[:40]
    return {"items": found, "folders": folders} if with_status else found


def _scan_manuscript_folder(docs_dir: str, home: str, ignored: set[str]) -> list[dict]:
    found = []
    for root, dirs, files in os.walk(docs_dir, onerror=lambda e: None):
        dirs[:] = [d for d in dirs if not d.startswith('.') and d.lower() not in ignored]
        
        for d in list(dirs):
            if d.lower().endswith('.scriv'):
                abs_p = os.path.join(root, d)
                found.append({
                    "type": "SCRIVENER",
                    "title": os.path.splitext(d)[0],
                    "path": abs_p,
                    "rel_path": os.path.relpath(abs_p, home)
                })
                dirs.remove(d)

        for f in files:
            if f.lower().endswith(('.docx', '.rtf')) and not f.startswith(('~$', '.')):
                abs_p = os.path.join(root, f)
                found.append({
                    "type": "RTF" if f.lower().endswith(".rtf") else "DOCX",
                    "title": os.path.splitext(f)[0].replace("_", " "),
                    "path": abs_p,
                    "rel_path": os.path.relpath(abs_p, home)
                })

    return found
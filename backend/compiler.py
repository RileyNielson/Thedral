import os
from backend.text_utils import html_to_formatted_paragraphs
import re
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from backend.database import get_db
from backend.config import EXPORTS_DIR

def compile_book_to_docx(
    book_id: str | None = None, 
    output_filename: str = "Thedral_Manuscript.docx",
    author_name: str = "",
    subtitle: str = "",
    dedication: str = "",
    chapter_ids: list[str] | None = None,
) -> str:
    """
    Compiles the active book into an industry-standard formatted Word manuscript (.docx).
    Features: 1-inch margins, Times New Roman, centered chapter headings,
    indented body paragraphs, centered scene dividers, and publication-grade epigraphs.
    """
    conn = get_db()
    c = conn.cursor()

    # 1. Resolve Active Book Node
    if book_id:
        c.execute(
            "SELECT id, title FROM binder_nodes WHERE id = ? AND node_type = 'BOOK'",
            (book_id,),
        )
        book = c.fetchone()
        if not book:
            conn.close()
            raise ValueError("The selected manuscript was not found.")
    else:
        c.execute("""
            SELECT id, title FROM binder_nodes 
            WHERE node_type = 'BOOK' AND (is_archived IS NULL OR is_archived = 0) 
            ORDER BY sort_order ASC LIMIT 1
        """)
        book = c.fetchone()

    if not book:
        conn.close()
        raise ValueError("No book found to compile in the database.")

    book_title = book["title"]
    doc = docx.Document()

    # 2. Configure 1-Inch Standard Margins
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    # 3. Build Title Page Header Block
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(72)
    title_p.paragraph_format.space_after = Pt(24)
    run_t = title_p.add_run(f"{book_title}\n")
    run_t.font.name = "Times New Roman"
    run_t.font.size = Pt(24)
    run_t.bold = True
    if subtitle.strip():
        subtitle_p = doc.add_paragraph()
        subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_r = subtitle_p.add_run(subtitle.strip())
        subtitle_r.font.name = "Times New Roman"
        subtitle_r.font.size = Pt(14)
        subtitle_r.italic = True
    if author_name.strip():
        author_p = doc.add_paragraph()
        author_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        author_p.paragraph_format.space_before = Pt(24)
        author_r = author_p.add_run(author_name.strip())
        author_r.font.name = "Times New Roman"
        author_r.font.size = Pt(12)
    if dedication.strip():
        doc.add_page_break()
        dedication_p = doc.add_paragraph()
        dedication_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        dedication_p.paragraph_format.space_before = Pt(216)
        dedication_r = dedication_p.add_run(dedication.strip())
        dedication_r.font.name = "Times New Roman"
        dedication_r.font.size = Pt(12)

    # 4. Fetch Chapters in Chronological Sequence
    chapter_query = """
        SELECT id, title, sort_order, epigraph
        FROM binder_nodes
        WHERE parent_id = ? AND node_type = 'CHAPTER'
    """
    chapter_params: list = [book["id"]]
    if chapter_ids is not None:
        if not chapter_ids:
            conn.close()
            raise ValueError("Select at least one chapter to export.")
        placeholders = ",".join("?" for _ in chapter_ids)
        chapter_query += f" AND id IN ({placeholders})"
        chapter_params.extend(chapter_ids)
    chapter_query += " ORDER BY sort_order ASC"
    c.execute(chapter_query, chapter_params)
    chapters = c.fetchall()
    if chapter_ids is not None and len(chapters) != len(set(chapter_ids)):
        conn.close()
        raise ValueError("One or more selected chapters do not belong to this manuscript.")

    for ch in chapters:
        doc.add_page_break()
        ch_p = doc.add_paragraph()
        ch_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ch_p.paragraph_format.space_before = Pt(72)
        ch_p.paragraph_format.space_after = Pt(18)
        ch_r = ch_p.add_run(ch["title"])
        ch_r.font.name = "Times New Roman"
        ch_r.font.size = Pt(16)
        ch_r.bold = True

        # Render Publication-Grade Decoupled Epigraph
        epigraph_text = ch["epigraph"] if "epigraph" in ch.keys() else ""
        if epigraph_text and epigraph_text.strip():
            epi_p = doc.add_paragraph()
            epi_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            epi_p.paragraph_format.left_indent = Inches(1.5)
            epi_p.paragraph_format.right_indent = Inches(1.5)
            epi_p.paragraph_format.space_before = Pt(12)
            epi_p.paragraph_format.space_after = Pt(24)
            epi_r = epi_p.add_run(f'"{epigraph_text.strip()}"')
            epi_r.font.name = "Times New Roman"
            epi_r.font.size = Pt(10.5)
            epi_r.italic = True

        # Fetch Scenes under this Chapter
        c.execute("""
            SELECT title, content, sort_order 
            FROM binder_nodes 
            WHERE parent_id = ? AND node_type = 'SCENE' 
            ORDER BY sort_order ASC
        """, (ch["id"],))
        scenes = c.fetchall()

        for idx, sc in enumerate(scenes):
            # Render Scene Break Dinkus between sibling scenes
            if idx > 0:
                dp = doc.add_paragraph()
                dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                dp.paragraph_format.first_line_indent = Inches(0)
                dp.paragraph_format.space_before = Pt(14)
                dp.paragraph_format.space_after = Pt(14)
                dr = dp.add_run("* * *")
                dr.font.name = "Times New Roman"
                dr.font.size = Pt(12)

            raw_content = sc["content"] or ""
            for para in html_to_formatted_paragraphs(raw_content):
                clean = "".join(text for text, _, _ in para).strip()
                if not clean:
                    continue

                # Catch manual scene break dividers in prose
                if clean in {"* * *", "###", "---", "~~~"}:
                    dp = doc.add_paragraph()
                    dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    dp.paragraph_format.first_line_indent = Inches(0)
                    dp.paragraph_format.space_before = Pt(14)
                    dp.paragraph_format.space_after = Pt(14)
                    dr = dp.add_run("* * *")
                    dr.font.name = "Times New Roman"
                    dr.font.size = Pt(12)
                    continue

                p_obj = doc.add_paragraph()
                p_obj.paragraph_format.first_line_indent = Inches(0.5)
                p_obj.paragraph_format.line_spacing = 1.5
                p_obj.paragraph_format.space_after = Pt(0)
                for text, bold, italic in para:
                    run = p_obj.add_run(text)
                    run.bold = bold
                    run.italic = italic
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(12)

    conn.close()

    # 5. Save Document into Verified Exports Directory
    os.makedirs(EXPORTS_DIR, exist_ok=True)
    out_path = os.path.join(EXPORTS_DIR, output_filename)
    doc.save(out_path)
    return out_path

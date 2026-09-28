import os
import time
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from core.story_db import get_db, WORKSPACE_DIR

EXPORTS_DIR = os.path.join(WORKSPACE_DIR, "Exports")
os.makedirs(EXPORTS_DIR, exist_ok=True)

def save_paragraph_edit(paragraph_id: str, new_text: str, change_note: str = "") -> int:
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT text, version, book_id FROM paragraphs WHERE id = ?", (paragraph_id,))
    row = c.fetchone()
    if not row:
        conn.close()
        return 0
        
    old_text = row["text"]
    current_ver = row["version"] or 1
    
    c.execute("""
    INSERT INTO paragraph_revisions (paragraph_id, version_num, text, change_note, created_at)
    VALUES (?, ?, ?, ?, ?)
    """, (paragraph_id, current_ver, old_text, change_note, time.strftime("%Y-%m-%d %H:%M:%S")))
    
    new_ver = current_ver + 1
    new_wc = len(new_text.split())
    c.execute("""
    UPDATE paragraphs SET text = ?, word_count = ?, version = ? WHERE id = ?
    """, (new_text, new_wc, new_ver, paragraph_id))
    
    conn.commit()
    conn.close()
    return new_ver

def rollback_paragraph(paragraph_id: str, target_version: int) -> bool:
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT text FROM paragraph_revisions WHERE paragraph_id = ? AND version_num = ?", 
              (paragraph_id, target_version))
    row = c.fetchone()
    if not row:
        conn.close()
        return False
        
    restored_text = row["text"]
    new_wc = len(restored_text.split())
    
    c.execute("UPDATE paragraphs SET text = ?, word_count = ?, version = ? WHERE id = ?",
              (restored_text, new_wc, target_version, paragraph_id))
    conn.commit()
    conn.close()
    return True

def get_revision_history(paragraph_id: str) -> list:
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT version_num, text, change_note, created_at FROM paragraph_revisions WHERE paragraph_id = ? ORDER BY version_num DESC", (paragraph_id,))
    rows = c.fetchall()
    conn.close()
    return rows

def check_downstream_ripples(book_id: str, from_chapter: int, from_para: int, keywords: list) -> list:
    if not keywords: return []
    conn = get_db()
    c = conn.cursor()
    c.execute("""
    SELECT id, chapter_num, para_num, text FROM paragraphs 
    WHERE book_id = ? AND (chapter_num > ? OR (chapter_num = ? AND para_num > ?))
    ORDER BY chapter_num ASC, para_num ASC
    """, (book_id, from_chapter, from_chapter, from_para))
    rows = c.fetchall()
    conn.close()
    
    ripples = []
    for r in rows:
        text_lower = r["text"].lower()
        matched = [k for k in keywords if k.lower() in text_lower]
        if matched:
            ripples.append({
                "id": r["id"],
                "chapter": r["chapter_num"],
                "para": r["para_num"],
                "matched_keywords": matched,
                "snippet": r["text"][:140] + "..."
            })
            if len(ripples) >= 8: break
                
    return ripples

def export_book_to_docx(book_id: str) -> str:
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT b.title, s.title as series_name, b.volume_num FROM books b JOIN series s ON b.series_id = s.id WHERE b.id = ?", (book_id,))
    b_info = c.fetchone()
    
    c.execute("SELECT chapter_num, para_num, text FROM paragraphs WHERE book_id = ? ORDER BY chapter_num ASC, para_num ASC", (book_id,))
    paras = c.fetchall()
    conn.close()
    
    if not paras: return ""

    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run(f"{b_info['title']}\n")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(24)
    run_title.bold = True
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run(f"{b_info['series_name']} — Volume {b_info['volume_num']}\n\n\n")
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(14)
    run_sub.italic = True

    current_ch = 0
    for p in paras:
        if p["chapter_num"] != current_ch:
            current_ch = p["chapter_num"]
            doc.add_page_break()
            ch_p = doc.add_paragraph()
            ch_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ch_run = ch_p.add_run(f"Chapter {current_ch}\n\n")
            ch_run.font.name = "Times New Roman"
            ch_run.font.size = Pt(16)
            ch_run.bold = True

        para_obj = doc.add_paragraph()
        para_obj.paragraph_format.first_line_indent = Inches(0.5)
        para_obj.paragraph_format.line_spacing = 1.5
        run = para_obj.add_run(p["text"])
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)

    filename = f"{b_info['series_name']}_Book{b_info['volume_num']:02d}_{b_info['title']}.docx".replace(" ", "_")
    output_path = os.path.join(EXPORTS_DIR, filename)
    doc.save(output_path)
    print(f"📄 [Exporter]: Saved to _Studio_Workspace/Exports/{filename}")
    return output_path

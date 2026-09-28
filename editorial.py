import os
import re
import time
import glob
import math
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from config import WORKSPACE_DIR, MIRROR_DIR, EXPORTS_DIR
from db import get_db

FILTER_VERBS = ["saw", "heard", "felt", "noticed", "watched", "wondered", "realized", "seemed", "appeared", "decided", "noted"]
CRUTCH_WORDS = ["suddenly", "glanced", "nodded", "frowned", "slightly", "almost", "began to", "just", "really", "very", "a bit", "shrugged"]
TENSION_KEYWORDS = ["blood", "blade", "death", "kill", "fire", "screamed", "gun", "shout", "danger", "run", "breath", "clutched", "strike", "shadow", "terror", "betrayal", "poison", "iron", "dark", "lethal", "collapse", "urgent"]

def calculate_chapter_telemetry(book_id: str, chapter_num: int) -> dict:
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT text, word_count FROM paragraphs WHERE book_id = ? AND chapter_num = ? ORDER BY seq_order ASC", (book_id, chapter_num))
    rows = c.fetchall()
    if not rows:
        conn.close()
        return {}
    total_words = sum([r["word_count"] for r in rows])
    total_paras = len(rows)
    all_text = " ".join([r["text"] for r in rows])
    text_lower = all_text.lower()
    filters_found = {v: len(re.findall(r"\b" + v + r"\b", text_lower)) for v in FILTER_VERBS if v in text_lower}
    crutches_found = {w: len(re.findall(r"\b" + w + r"\b", text_lower)) for w in CRUTCH_WORDS if w in text_lower}
    sentences = [s.strip() for s in re.split(r'[.!?]+', all_text) if s.strip()]
    sentence_lengths = [len(s.split()) for s in sentences if len(s.split()) > 0]
    if sentence_lengths:
        avg_len = round(sum(sentence_lengths) / len(sentence_lengths), 1)
        variance = round(sum((x - avg_len) ** 2 for x in sentence_lengths) / len(sentence_lengths), 1)
        stdev = round(math.sqrt(variance), 1)
    else: avg_len, variance, stdev = 0.0, 0.0, 0.0
    dialogue_paras = sum([1 for r in rows if '"' in r["text"] or '“' in r["text"]])
    dialogue_ratio = round((dialogue_paras / max(1, total_paras)) * 100, 1)
    tension_hits = sum(len(re.findall(r"\b" + k + r"\b", text_lower)) for k in TENSION_KEYWORDS)
    tension_density = (tension_hits / max(1, total_words)) * 1000
    elev_score = 0.35 + (0.35 * (dialogue_ratio / 100.0)) + min(0.25, tension_density * 0.02)
    if avg_len < 12: elev_score += 0.10
    elif avg_len > 22: elev_score -= 0.10
    tension_elevation = round(max(0.05, min(0.95, elev_score)), 2)
    pacing_slope = 0.0
    slope_diagnosis = "STABLE_FLOW"
    if chapter_num > 1:
        c.execute("SELECT dialogue_ratio FROM chapter_digests WHERE book_id = ? AND chapter_num = ?", (book_id, chapter_num - 1))
        prev_digest = c.fetchone()
        if prev_digest:
            prev_elev = 0.35 + (0.35 * prev_digest["dialogue_ratio"])
            pacing_slope = round(tension_elevation - prev_elev, 2)
            if abs(pacing_slope) < 0.04: slope_diagnosis = "EXPOSITION_PLATEAU"
            elif pacing_slope > 0.45: slope_diagnosis = "MELODRAMATIC_WHIPLASH"
            elif pacing_slope > 0.15: slope_diagnosis = "RISING_MOMENTUM"
            elif pacing_slope < -0.15: slope_diagnosis = "CATHARTIC_DESCENT"
    conn.close()
    return {
        "chapter_num": chapter_num, "total_words": total_words, "total_paras": total_paras,
        "avg_sentence_len": avg_len, "cadence_variance": variance, "cadence_stdev": stdev,
        "dialogue_ratio": dialogue_ratio, "tension_elevation": tension_elevation,
        "pacing_velocity": stdev, "pacing_slope": pacing_slope, "slope_diagnosis": slope_diagnosis,
        "top_filters": dict(sorted(filters_found.items(), key=lambda i: i[1], reverse=True)[:5]),
        "top_crutches": dict(sorted(crutches_found.items(), key=lambda i: i[1], reverse=True)[:5])
    }

def export_local_mirror_files(book_id: str):
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT DISTINCT chapter_num FROM paragraphs WHERE book_id = ? ORDER BY chapter_num ASC", (book_id,))
    chapters = [r[0] for r in c.fetchall()]
    for ch in chapters:
        c.execute("SELECT id, text FROM paragraphs WHERE book_id = ? AND chapter_num = ? ORDER BY seq_order ASC", (book_id, ch))
        rows = c.fetchall()
        target_path = os.path.join(MIRROR_DIR, f"Chapter_{ch:02d}.md")
        temp_path = target_path + ".tmp"
        with open(temp_path, "w", encoding="utf-8") as f:
            f.write(f"# Chapter {ch}\n\n")
            for r in rows: f.write(f"<!-- ID: {r['id']} -->\n{r['text']}\n\n")
        os.replace(temp_path, target_path)
    conn.close()

def sync_from_local_mirror(book_id: str) -> int:
    conn = get_db()
    updated_count = 0
    now_ts = time.strftime("%Y-%m-%d %H:%M:%S")
    with conn:
        c = conn.cursor()
        for ch_file in sorted(glob.glob(os.path.join(MIRROR_DIR, "Chapter_*.md"))):
            with open(ch_file, "r", encoding="utf-8") as f: content = f.read()
            blocks = re.split(r"<!-- ID:\s*(\w+)\s*-->", content)
            for i in range(1, len(blocks), 2):
                pid, body = blocks[i].strip(), blocks[i+1].strip()
                c.execute("SELECT para_uuid, text, version FROM paragraphs WHERE id = ? AND book_id = ?", (pid, book_id))
                row = c.fetchone()
                if row and row["text"] != body:
                    old_ver = row["version"] or 1
                    c.execute("INSERT INTO paragraph_revisions (para_uuid, display_id, version_num, text, change_note, created_at) VALUES (?, ?, ?, ?, 'Synced from Mac mirror', ?)",
                              (row["para_uuid"], pid, old_ver, row["text"], now_ts))
                    c.execute("UPDATE paragraphs SET text = ?, word_count = ?, version = ?, updated_at = ? WHERE id = ? AND book_id = ?",
                              (body, len(body.split()), old_ver + 1, now_ts, pid, book_id))
                    updated_count += 1
    conn.close()
    return updated_count

def check_downstream_ripples(book_id: str, from_chapter: int, from_para: int, keywords: list) -> list:
    stop_words = {"the", "and", "with", "that", "this", "from", "they", "their"}
    clean_kw = [re.escape(k.strip('.,;:"!?')) for k in keywords if len(k) > 3 and k.lower() not in stop_words]
    if not clean_kw: return []
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT id, chapter_num, para_num, text FROM paragraphs WHERE book_id = ? AND (chapter_num > ? OR (chapter_num = ? AND para_num > ?)) ORDER BY seq_order ASC",
              (book_id, from_chapter, from_chapter, from_para))
    rows = c.fetchall()
    conn.close()
    ripples = []
    pattern = re.compile(r"\b(" + "|".join(clean_kw) + r")\b", re.IGNORECASE)
    for r in rows:
        matches = pattern.findall(r["text"])
        if matches:
            ripples.append({
                "id": r["id"], "chapter": r["chapter_num"], "para": r["para_num"],
                "matched": sorted(list(set(m.title() for m in matches))),
                "snippet": r["text"][:140] + ("..." if len(r["text"]) > 140 else "")
            })
            if len(ripples) >= 5: break
    return ripples

def export_book_to_docx(book_id: str) -> str:
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT title FROM books WHERE id = ?", (book_id,))
    b_info = c.fetchone()
    c.execute("SELECT chapter_num, text FROM paragraphs WHERE book_id = ? ORDER BY seq_order ASC", (book_id,))
    paras = c.fetchall()
    conn.close()
    if not paras: return ""
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Inches(1), Inches(1), Inches(1), Inches(1)
    book_title = b_info['title'] if b_info else 'Manuscript'
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before, title_p.paragraph_format.space_after = Pt(72), Pt(24)
    run_t = title_p.add_run(f"{book_title}\n")
    run_t.font.name, run_t.font.size, run_t.bold = "Times New Roman", Pt(24), True
    current_ch = 0
    for p in paras:
        text, ch_num = p["text"].strip(), p["chapter_num"]
        if ch_num != current_ch:
            current_ch = ch_num
            doc.add_page_break()
            ch_p = doc.add_paragraph()
            ch_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ch_p.paragraph_format.space_before, ch_p.paragraph_format.space_after = Pt(72), Pt(24)
            ch_r = ch_p.add_run(f"Chapter {current_ch}")
            ch_r.font.name, ch_r.font.size, ch_r.bold = "Times New Roman", Pt(16), True
        if text in ["* * *", "###", "---"]:
            dp = doc.add_paragraph()
            dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            dp.paragraph_format.space_before, dp.paragraph_format.space_after = Pt(12), Pt(12)
            dr = dp.add_run("* * *")
            dr.font.name, dr.font.size = "Times New Roman", Pt(12)
            continue
        para_obj = doc.add_paragraph()
        para_obj.paragraph_format.first_line_indent, para_obj.paragraph_format.line_spacing = Inches(0.5), 1.5
        run = para_obj.add_run(text)
        run.font.name, run.font.size = "Times New Roman", Pt(12)
    safe_title = re.sub(r'[\W_]+', '_', book_title)
    output_path = os.path.join(EXPORTS_DIR, f"{safe_title}_Standard_Export.docx")
    doc.save(output_path)
    return output_path

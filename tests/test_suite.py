"""
Thedral Sovereign Studio — Master Architecture Verification Suite
Executes end-to-end integration tests across:
1. SQLite WAL & FTS5 Synchronization Triggers
2. Span-Linked Semantic Hashing & Severed Thread Detection
3. Law 1 Pacing Physics & Sentence Energetics (Craft X-Ray)
4. Clutter Shield & Entity Taxonomy Guardrails
5. 8-Layer 3D Cosmograph & Beat-Level Chapter Lens
6. Scrivener & Word Importer Paragraph Normalization
7. Voice Profile & Modern Anachronism Detection
8. FastAPI Router API Contracts & Endpoints
9. Transparent text-pattern signal analysis
"""

import sys
import os
import re
import json
import sqlite3
import tempfile
import types
import uuid
from unittest.mock import patch
from pathlib import Path

# Add project root and current directory to Python sys.path
current_dir = os.path.abspath(os.path.dirname(__file__))
project_root = os.path.abspath(os.path.join(current_dir, ".."))
for p in [project_root, current_dir]:
    if p not in sys.path:
        sys.path.insert(0, p)

# Terminal formatting
PASS = "\033[92m[PASS]\033[0m"
FAIL = "\033[91m[FAIL]\033[0m"
HEADER = "\033[95m"
BOLD = "\033[1m"
RESET = "\033[0m"

tests_run = 0
tests_passed = 0

def record_test(name: str, passed: bool, detail: str = ""):
    global tests_run, tests_passed
    tests_run += 1
    if passed:
        tests_passed += 1
        print(f" {PASS} {name}")
    else:
        print(f" {FAIL} {name} — {detail}")

print(f"\n{BOLD}{HEADER}==============================================================================={RESET}")
print(f"{BOLD}{HEADER}       THEDRAL MASTER ARCHITECTURE VERIFICATION MEGA-SUITE                      {RESET}")
print(f"{BOLD}{HEADER}==============================================================================={RESET}\n")

# =============================================================================
# 1. DATABASE & FTS5 INTEGRATION
# =============================================================================
print(f"{BOLD}1. Database Engine, Migrations, & FTS5 Instant Search{RESET}")

with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as tmp:
    test_db_path = tmp.name

try:
    from backend.models import init_db
    from backend.database import get_db

    init_db(test_db_path)
    conn = get_db(test_db_path)
    c = conn.cursor()

    c.execute("SELECT name FROM sqlite_master WHERE type='table'")
    tables = [r[0] for r in c.fetchall()]
    required_tables = [
        "binder_nodes", "binder_fts", "canonical_entities", "canonical_worldlines",
        "entity_relationships", "style_sheet", "node_revisions", "promises_ledger",
        "publishing_profiles", "export_history", "submission_records"
    ]
    all_tables_present = all(t in tables for t in required_tables)
    record_test("Database Schema & Self-Healing Migrations", all_tables_present)

    c.execute("""
        INSERT INTO binder_nodes (id, parent_id, node_type, title, content, sort_order)
        VALUES ('test_scene_1', NULL, 'SCENE', 'The Iron Balcony', 'Vance checked the bronze pressure gauge.', 10.0)
    """)
    conn.commit()

    c.execute("SELECT node_id FROM binder_fts WHERE binder_fts MATCH 'gauge'")
    fts_match = c.fetchone()
    record_test("SQLite FTS5 Automated Insertion Trigger", fts_match is not None and fts_match[0] == 'test_scene_1')

    c.execute("UPDATE binder_nodes SET content = 'The clockwork automaton whirred.' WHERE id = 'test_scene_1'")
    conn.commit()
    c.execute("SELECT node_id FROM binder_fts WHERE binder_fts MATCH 'automaton'")
    fts_update_match = c.fetchone()
    c.execute("SELECT node_id FROM binder_fts WHERE binder_fts MATCH 'gauge'")
    old_gauge_match = c.fetchone()
    record_test("SQLite FTS5 Automated Update Trigger (Old Term Purged)", fts_update_match is not None and old_gauge_match is None)

    conn.close()
finally:
    if os.path.exists(test_db_path):
        os.remove(test_db_path)

# =============================================================================
# 2. SPAN-LINKED CAUSALITY & SEVERED THREADS
# =============================================================================
print(f"\n{BOLD}2. Span-Linked Semantic Hashing & Severed Thread Detection{RESET}")

from backend.hashing import compute_block_hash, get_scene_block_hashes, find_severed_threads

p1 = "Bare marble met my soles like frozen iron. My lungs locked midway through an exhale."
p2 = "Captain Vance dropped the brass cylinder onto the desk."

h1 = compute_block_hash(p1)
h2 = compute_block_hash(p2)

p1_typo = "bare marble met my soles like frozen iron! my lungs locked midway through an exhale"
h1_typo = compute_block_hash(p1_typo)

record_test("Semantic Block Hashing Generates 16-Char Hex", len(h1) == 16 and len(h2) == 16)
record_test("Punctuation & Case Normalized Invariance", h1 == h1_typo)

test_conn = sqlite3.connect(":memory:")
test_conn.row_factory = sqlite3.Row
test_conn.execute("""
    CREATE TABLE canonical_worldlines (
        event_id TEXT PRIMARY KEY,
        entity_id TEXT,
        timeline_tick REAL,
        scene_id TEXT,
        source_span_hash TEXT,
        state_delta TEXT,
        epistemic_tier TEXT
    )
""")
test_conn.execute("""
    INSERT INTO canonical_worldlines VALUES ('ev_1', 'vance', 1.0, 'sc_1', ?, '{"health": 0.5}', 'FACT')
""", (h2,))

severed_before = find_severed_threads(test_conn, 'sc_1', f"{p1}\n\n{p2}")
record_test("Intact Paragraph Generates Zero Severed Threads", len(severed_before) == 0)

severed_after = find_severed_threads(test_conn, 'sc_1', p1)
record_test("Excising Paragraph Detects Severed Causal Thread", len(severed_after) == 1 and severed_after[0]["entity_id"] == "vance")
test_conn.close()

# =============================================================================
# 3. CRAFT X-RAY & PACING TELEMETRY (WITH IDIOM SHIELD)
# =============================================================================
print(f"\n{BOLD}3. Law 1 Pacing Physics, Craft X-Ray, & Idiom Shield{RESET}")

from backend.telemetry import calculate_text_telemetry, classify_sentence_energetics, scrub_idioms

sample_prose = (
    "The outer courtyard bells broke through the clerestory vents. "
    "Bare marble met my soles like frozen iron. "
    "My lungs locked midway through an exhale! "
    "\"Where is the ledger?\" Vance whispered. "
    "I saw that he appeared to decide to watch the door slightly frowned."
)

telemetry = calculate_text_telemetry(sample_prose)
record_test("Cadence Velocity & Sentence Standard Deviation", telemetry["cadence_stdev"] > 0)
record_test("Dialogue Ratio Calculation", telemetry["dialogue_ratio"] > 0)
record_test("Filter Verbs & Crutch Words Scanning", len(telemetry["filters"]) >= 2 and len(telemetry["crutches"]) >= 1)

# Anti-GIGO Idiom Shield test
idiom_text = "She drank blood-orange tea by the fireplace, bored to death."
scrubbed = scrub_idioms(idiom_text)
record_test("Anti-GIGO Idiom Shield Scrubs False Crises (Blood-Orange/Fireplace)", "blood" not in scrubbed.lower() and "fire" not in scrubbed.lower())

xray = classify_sentence_energetics(sample_prose)
types_found = {s["type"] for s in xray}
record_test("Craft X-Ray Classifies Energetics (Escalators/Slack/Pivots)", "ESCALATOR" in types_found and "SLACK" in types_found)

# =============================================================================
# 4. CLUTTER SHIELD & ENTITY TAXONOMY
# =============================================================================
print(f"\n{BOLD}4. Clutter Shield & Entity Verification{RESET}")

from backend.guardrails import is_valid_character_name, sanitize_entity_type
from backend.entity_names import entity_name_forms, parse_entity_aliases, text_mentions_entity

record_test("Valid Character Allowed (Ellie)", is_valid_character_name("Ellie"))
record_test("Valid Character Allowed (Lord Malakor)", is_valid_character_name("Lord Malakor"))
record_test("Clutter Shield Blocks Common Nouns (Meat Pie)", not is_valid_character_name("Meat Pie"))
record_test("Clutter Shield Blocks Furniture (Crutches)", not is_valid_character_name("Crutches"))
record_test("Clutter Shield Blocks Generic Mobs (Guards)", not is_valid_character_name("The Guards"))

record_test("Sanitize Demotes Clutter to CLUTTER", sanitize_entity_type("Lantern") == "CLUTTER")
record_test("Sanitize Preserves Legitimate Character", sanitize_entity_type("Captain Vance", "CHARACTER") == "CHARACTER")
record_test("Author-confirmed worldbuilding terms remain trackable", sanitize_entity_type("Aetherite", "WORLD_TERM") == "WORLD_TERM")
record_test(
    "Alias parsing supports old delimiters and new JSON name lists",
    parse_entity_aliases('["Lady Mara", "Mare", "Captain Vale"]')
    == ["Lady Mara", "Mare", "Captain Vale"]
    and parse_entity_aliases("Lady Mara, Mare; Captain Vale")
    == ["Lady Mara", "Mare", "Captain Vale"],
)
record_test(
    "Alias matching is case-insensitive, whole-name, and deduplicated",
    entity_name_forms("Mara Vale", '["Mare", "MARA VALE"]') == ["Mara Vale", "Mare"]
    and text_mentions_entity("The Captain Vale entered.", "Mara Vale", '["Captain Vale"]')
    and not text_mentions_entity("The Captain Valerian waited.", "Mara Vale", '["Vale"]'),
)

from backend.routers.entities import _candidate_names
candidate_scenes = [
    {"id": "scene-a", "title": "A", "content": "Rowan met Aetherite beneath the Silver Archive."},
    {"id": "scene-b", "title": "B", "content": "They found Aetherite glowing beneath the Silver Archive."},
]
candidate_results = _candidate_names(candidate_scenes, [])
record_test(
    "Entity discovery surfaces recurring invented terms with evidence",
    {candidate["name"] for candidate in candidate_results} == {"Aetherite", "Silver Archive"}
    and all(candidate["mentions"] == 2 and candidate["examples"] for candidate in candidate_results),
)
registered_candidates = _candidate_names(
    candidate_scenes,
    [{"name": "Aetherite", "aliases": "The Aether"}],
)
record_test(
    "Entity discovery excludes names already confirmed in canon",
    all(candidate["name"] != "Aetherite" for candidate in registered_candidates),
)

# =============================================================================
# 5. 3D ASTROLABE (COSMOGRAPH & CHAPTER LENS)
# =============================================================================
print(f"\n{BOLD}5. 8-Layer Spacetime Cosmograph & Beat-Level Chapter Lens{RESET}")

from backend.astrolabe import build_astrolabe_manifest

astro_conn = sqlite3.connect(":memory:")
astro_conn.row_factory = sqlite3.Row
astro_conn.execute("""
    CREATE TABLE binder_nodes (
        id TEXT PRIMARY KEY, parent_id TEXT, project_id TEXT DEFAULT 'default',
        node_type TEXT, title TEXT, sort_order REAL, content TEXT, word_count INTEGER,
        is_archived INTEGER DEFAULT 0, synopsis TEXT DEFAULT ''
    )
""")
astro_conn.execute("""
    CREATE TABLE entity_relationships (source_id TEXT, target_id TEXT, rel_type TEXT, description TEXT)
""")
astro_conn.execute("""
    CREATE TABLE canonical_worldlines (
        event_id TEXT PRIMARY KEY, entity_id TEXT, timeline_tick REAL, scene_id TEXT,
        source_span_hash TEXT, state_delta TEXT, epistemic_tier TEXT, carrier_id TEXT
    )
""")
astro_conn.execute("""
    CREATE TABLE canonical_entities (entity_id TEXT PRIMARY KEY, name TEXT, entity_type TEXT, aliases TEXT, created_at TEXT)
""")
astro_conn.execute("""
    CREATE TABLE promises_ledger (id INTEGER PRIMARY KEY, series_id TEXT, book_id TEXT, planted_scene_id TEXT, target_volume INTEGER, target_tick REAL, target_location TEXT, promise_desc TEXT, status TEXT)
""")
astro_conn.execute("""
    CREATE TABLE story_deadlines (id TEXT PRIMARY KEY, book_id TEXT, title TEXT, deadline_tick REAL, urgency_level TEXT, status TEXT)
""")

astro_conn.execute("INSERT INTO binder_nodes VALUES ('book_1', NULL, 'default', 'BOOK', 'Book One', 10.0, '', 0, 0, '')")
astro_conn.execute("INSERT INTO binder_nodes VALUES ('ch_1', 'book_1', 'default', 'CHAPTER', 'Chapter One', 10.0, '', 0, 0, '')")
three_beat_prose = (
    "She considered the Docks and wondered what had happened to the light.\n\n"
    "\"Drop the weapon!\" the guard shouted near the North Gate and Gatehouse.\n\n"
    "Cold steel cleaved through the night air. Valentine watched; Vale ran past the Gatehouse."
)
astro_conn.execute("""
    INSERT INTO binder_nodes VALUES ('sc_1', 'ch_1', 'default', 'SCENE', 'The Confrontation', 10.0, ?, 36, 0, 'Combat at Docks')
""", (three_beat_prose,))
astro_conn.execute("""
    INSERT INTO canonical_entities VALUES
    ('mara', 'Mara Vale', 'CHARACTER', 'Captain Vale, Vale', ''),
    ('iron_docks', 'Iron Docks', 'LOCATION', 'the Docks', ''),
    ('north_gate', 'North Gate', 'LOCATION', 'Gate', '')
""")
astro_conn.commit()

macro_manifest = build_astrolabe_manifest(astro_conn, book_id='book_1')
record_test("Macro Cosmograph Generates 3D Topography Traces", len(macro_manifest["data"]) >= 1)

micro_manifest = build_astrolabe_manifest(astro_conn, chapter_id='ch_1', focus_scene_id='sc_1')
beats_trace = micro_manifest["data"][0] if micro_manifest["data"] else {}
num_beats = len(beats_trace.get("x", []))
record_test("Chapter Lens Maps Paragraph Beats (Not 1 Dot)", num_beats == 3)
record_test("Chapter Lens Encodes Clickable Beat Navigation Data", len(beats_trace.get("customdata", [])) == 3 and isinstance(beats_trace["customdata"][0], dict))
micro_cast = next((trace for trace in micro_manifest["data"] if trace.get("name") == "Cast: Mara Vale"), {})
record_test(
    "Chapter Lens Matches Confirmed Character Aliases Without Substring False Positives",
    len(micro_cast.get("x", [])) == 1,
)
macro_mention = next((trace for trace in macro_manifest["data"] if trace.get("name") == "Text mentions: Mara Vale"), {})
record_test(
    "Cosmograph Maps Text Mentions and Confirmed Location Aliases",
    len(macro_mention.get("x", [])) == 1
    and macro_mention["text"][0].endswith("Iron Docks"),
)

astro_conn.close()

# =============================================================================
# 6. SCRIVENER & TEXT IMPORTER NORMALIZATION
# =============================================================================
print(f"\n{BOLD}6. Importer Paragraph Normalization & Scene Assembly{RESET}")

from backend.importer import (
    normalize_paragraphs, strip_rtf_native, parse_docx, build_binder_from_manifest,
    parse_scrivener_project, build_import_report, inspect_docx_import
)
from backend.text_utils import html_to_formatted_paragraphs, html_to_text

single_newline_rtf = "Line one of paragraph.\nLine two of paragraph that continues.\nAnother distinct paragraph entirely."
paras = normalize_paragraphs(single_newline_rtf)
record_test("Normalizes textutil Single-Newline Streams", len(paras) == 3)
rtf_text = strip_rtf_native(r"{\rtf1\ansi Caf\'e9 \u8217?{ \b bold} prose\par next line}")
record_test("Decodes Scrivener RTF accents, Unicode, and paragraphs", rtf_text == "Café ’ bold prose\nnext line")
formatted_rtf = strip_rtf_native(r"{\rtf1 A {\b bold} and {\i italic}.}", preserve_formatting=True)
record_test("Preserves RTF bold and italic runs", formatted_rtf == "A <strong>bold</strong> and <em>italic</em>.")
formatted_html = "<p>A <strong>bold</strong> line.</p><p>An <em>italic</em> sentence.</p>"
formatted_paragraphs = html_to_formatted_paragraphs(formatted_html)
record_test(
    "Rich text stays readable to analysis and word counts",
    html_to_text(formatted_html) == "A bold line.\n\nAn italic sentence."
    and len(html_to_text(formatted_html).split()) == 6
    and len(formatted_paragraphs) == 2
    and any(text == "bold" and bold for text, bold, _ in formatted_paragraphs[0])
    and any(text == "italic" and italic for text, _, italic in formatted_paragraphs[1])
)

with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp_docx:
    formatted_docx_path = tmp_docx.name
try:
    from docx import Document
    formatted_doc = Document()
    paragraph = formatted_doc.add_paragraph()
    paragraph.add_run("Bold").bold = True
    paragraph.add_run(" and ")
    paragraph.add_run("italic").italic = True
    formatted_doc.save(formatted_docx_path)
    docx_manifest = parse_docx(formatted_docx_path)
    formatted_doc.add_table(rows=1, cols=1).cell(0, 0).text = "Tabular content"
    formatted_doc.save(formatted_docx_path)
    docx_warnings = inspect_docx_import(formatted_docx_path)
    record_test(
        "Preserves DOCX bold and italic runs",
        docx_manifest == [("<strong>Bold</strong> and <em>italic</em>", False, "", "", "")]
    )
    record_test("Reports unsupported DOCX tables", any("table" in warning.lower() for warning in docx_warnings))
finally:
    os.remove(formatted_docx_path)

preview_report = build_import_report({
    "manuscript": [
        ("Chapter: Act I", True, "Act I", "Chapter summary", ""),
        ("### SCENE_TITLE: Opening", False, "Opening", "", "Scene notes"),
        ("<p><strong>Bold</strong> opening.</p>", False, "Opening", "", ""),
        ("### SCENE_TITLE: Outline card", False, "Outline card", "A scene synopsis.", ""),
    ],
    "characters": [{"name": "Ellie"}],
    "locations": [{"name": "Dock"}],
    "lore": [{"term": "Aether"}],
})
record_test(
    "Import report summarizes scenes, prose, formatting, notes, and outline cards",
    preview_report["chapters"] == 1
    and preview_report["scenes"] == 2
    and preview_report["outline_only"] == 1
    and preview_report["paragraphs"] == 1
    and preview_report["words"] == 2
    and preview_report["formatted_paragraphs"] == 1
    and preview_report["notes"] == 1
    and preview_report["synopses"] == 2
    and preview_report["characters"] == 1
)

mock_manifest = {
    "manuscript": [
        ("Chapter: Act I", True, "Act I", "", ""),
        ("### SCENE_TITLE: The Infiltration", False, "The Infiltration", "Ellie slips in.", "Keep pacing tight."),
        ("Beat one prose.", False, "The Infiltration", "", ""),
        ("Beat two prose.", False, "The Infiltration", "", "")
    ],
    "characters": [{"name": "Ellie", "bio": "Master of locks." + ("A" * 2100), "synopsis": "Tallow smell"}],
    "locations": [{"name": "Aethelgard Docks", "desc": "Salt and cold iron."}],
    "lore": [{"term": "Aetherite", "definition": "Volatile blue crystal." + ("L" * 900), "metadata": {"uuid": "lore-a"}}]
}

with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as tmp_import:
    import_db_path = tmp_import.name

try:
    from backend.models import init_db
    init_db(import_db_path)

    import backend.importing.binder_builder
    backend.importing.binder_builder.get_db = lambda: sqlite3.connect(import_db_path)

    first_scene_id = build_binder_from_manifest(mock_manifest, "Test Import Project")

    verify_conn = sqlite3.connect(import_db_path)
    verify_conn.row_factory = sqlite3.Row
    c = verify_conn.cursor()

    c.execute("SELECT synopsis, notes, content FROM binder_nodes WHERE id = ?", (first_scene_id,))
    sc_row = c.fetchone()

    c.execute("SELECT name, axioms FROM canonical_entities WHERE entity_id = 'ellie'")
    ent_row = c.fetchone()

    c.execute("SELECT term, rule_definition, metadata_json FROM style_sheet WHERE term = 'Aetherite'")
    lore_row = c.fetchone()

    record_test("Importer Commits Scenes with Preserved Synopses & Notes", sc_row is not None and sc_row["synopsis"] == "Ellie slips in." and "Beat two prose." in sc_row["content"])
    record_test("Importer Extracts Full Character Sheet to Cast Directory", ent_row is not None and ent_row["name"] == "Ellie" and len(json.loads(ent_row["axioms"])["bio"]) == 2116)
    record_test("Importer Preserves Full Lore and Its Metadata", lore_row is not None and lore_row["term"] == "Aetherite" and len(lore_row["rule_definition"]) == 922 and json.loads(lore_row["metadata_json"])["scrivener"]["uuid"] == "lore-a")

    verify_conn.close()

    with tempfile.TemporaryDirectory(suffix=".scriv") as scriv_dir:
        scriv_root = Path(scriv_dir)
        (scriv_root / "Sample.scrivx").write_text(
            """<?xml version="1.0" encoding="UTF-8"?>
<ScrivenerProject Version="2.0"><Binder><BinderItem Type="Folder" UUID="draft"><Title>Draft</Title><Children>
  <BinderItem Type="Folder" UUID="chapter"><Title>Chapter One</Title><Synopsis>Chapter plan.</Synopsis><MetaData><MetaDataItem Key="Label">Blue</MetaDataItem></MetaData><Children>
    <BinderItem Type="Text" UUID="scene-a"><Title>Opening</Title><MetaData><MetaDataItem Key="Status">Revised</MetaDataItem><MetaDataItem Key="Keyword">storm</MetaDataItem><CustomField Name="POV">Mara</CustomField></MetaData><Children/></BinderItem>
    <BinderItem Type="Text" UUID="scene-b"><Title>Opening</Title><Children/></BinderItem>
    <BinderItem Type="Text" UUID="outline-only"><Title>Unwritten Turn</Title><Synopsis>A discovery changes the plan.</Synopsis><Notes>Keep the reveal quiet.</Notes><Children/></BinderItem>
  </Children></BinderItem>
</Children></BinderItem></Binder></ScrivenerProject>""",
            encoding="utf-8"
        )
        data_dir = scriv_root / "Files" / "Data"
        for doc_id, content, synopsis, notes in [
            ("scene-a", r"{\rtf1 First {\i scene} {\b content}.}", "Opening synopsis.", "Opening notes."),
            ("scene-b", "Second scene content.", "", ""),
            ("outline-only", "", "A discovery changes the plan.", "Keep the reveal quiet."),
        ]:
            doc_dir = data_dir / doc_id
            doc_dir.mkdir(parents=True)
            if content:
                filename = "content.rtf" if content.startswith(r"{\rtf") else "content.txt"
                (doc_dir / filename).write_text(content, encoding="utf-8")
            if synopsis:
                (doc_dir / "synopsis.txt").write_text(synopsis, encoding="utf-8")
            if notes:
                (doc_dir / "notes.txt").write_text(notes, encoding="utf-8")
        (data_dir / "scene-a" / "annotations.txt").write_text("Author annotation preserved.", encoding="utf-8")

        scriv_manifest = parse_scrivener_project(scriv_dir)
        scriv_first_scene = build_binder_from_manifest(scriv_manifest, "Sample")
        verify_conn = sqlite3.connect(import_db_path)
        verify_conn.row_factory = sqlite3.Row
        scenes = verify_conn.execute(
            "SELECT id, title, synopsis, notes, content, card_data FROM binder_nodes WHERE node_type = 'SCENE' ORDER BY rowid DESC LIMIT 3"
        ).fetchall()
        scene_data = [dict(row) for row in scenes]
        contents = {html_to_text(row["content"]): row for row in scenes}
        first_scene = contents.get("First scene content.")
        second_scene = contents.get("Second scene content.")
        outline_scene = next((row for row in scenes if row["title"] == "Unwritten Turn"), None)
        first_metadata = json.loads(first_scene["card_data"])["scrivener"]["metadata"] if first_scene else {}
        first_external_files = json.loads(first_scene["card_data"])["scrivener"].get("external_files", {}) if first_scene else {}
        outline_metadata = json.loads(outline_scene["card_data"])["scrivener"] if outline_scene else {}
        chapter_metadata = verify_conn.execute(
            "SELECT card_data FROM binder_nodes WHERE node_type = 'CHAPTER' ORDER BY rowid DESC LIMIT 1"
        ).fetchone()
        project_metadata = verify_conn.execute(
            "SELECT card_data FROM binder_nodes WHERE node_type = 'BOOK' ORDER BY rowid DESC LIMIT 1"
        ).fetchone()
        saved_project_metadata = json.loads(project_metadata["card_data"])["scrivener"] if project_metadata else {}
        record_test(
            "Scrivener maps content and notes by UUID without title collisions",
            len(scene_data) == 3 and first_scene is not None and first_scene["synopsis"] == "Opening synopsis."
            and first_scene["notes"] == "Opening notes." and second_scene is not None
            and "<em>scene</em>" in first_scene["content"] and "<strong>content</strong>" in first_scene["content"]
        )
        record_test(
            "Scrivener preserves outline-only documents and item metadata",
            outline_scene is not None and outline_scene["synopsis"] == "A discovery changes the plan."
            and outline_scene["notes"] == "Keep the reveal quiet."
            and first_metadata.get("MetaData") and outline_metadata.get("uuid") == "outline-only"
            and first_external_files.get("annotations") == "Author annotation preserved."
        )
        record_test(
            "Scrivener preserves chapter and project metadata",
            chapter_metadata is not None and "scrivener" in json.loads(chapter_metadata["card_data"])
            and verify_conn.execute("SELECT synopsis FROM binder_nodes WHERE id = (SELECT parent_id FROM binder_nodes WHERE id = ?)", (scriv_first_scene,)).fetchone()["synopsis"] == "Chapter plan."
            and saved_project_metadata.get("scrivx_file") == "Sample.scrivx"
            and any(folder["title"] == "Draft" for folder in saved_project_metadata.get("binder_folders", []))
            and scriv_first_scene in {row["id"] for row in scenes}
        )
        verify_conn.close()
finally:
    if os.path.exists(import_db_path):
        os.remove(import_db_path)

# =============================================================================
# 7. VOICE PROFILE & ANACHRONISM AUDITOR
# =============================================================================
print(f"\n{BOLD}7. Voice Profile & Anachronism Detection{RESET}")

from backend.voice import audit_voice_continuity

modern_text = "Ellie realized this toxic relationship was a red flag and had no bandwidth to circle back."
voice_audit = audit_voice_continuity(modern_text, pov_entity_id="ellie")
record_test("Detects Modern Anachronisms (Red Flag, Bandwidth, Toxic)", len(voice_audit["anachronisms"]) >= 3)
record_test("Deducts Voice Match Score for Dialect Contamination", voice_audit["voice_match"] < 80)

# =============================================================================
# 8. FASTAPI REST API CONTRACTS
# =============================================================================
print(f"\n{BOLD}8. FastAPI Router API Contracts & Endpoints{RESET}")

from backend.reader_story import analyze_reader_story
from backend.reader_story import _turns

story_read = analyze_reader_story(
    [
        {
            "scene_id": "setup",
            "title": "Before Dawn",
            "content": "Mara promised to find the silver key before dawn. She loved the dog and waited for him.",
        },
        {
            "scene_id": "payoff",
            "title": "The Floodgate",
            "content": (
                "Mara found the silver key, but the floodgate failed. "
                "The dog died before morning. Mara trembled alone. "
                "Mara asked, 'What happened?' She said nothing. Mara lied about the key. "
                "Mara said the passage was perfectly safe, then the bridge failed."
            ),
        },
    ],
    "payoff",
    (
        "Mara found the silver key, but the floodgate failed. "
        "The dog died before morning. Mara trembled alone. "
        "Mara asked, 'What happened?' She said nothing. Mara lied about the key. "
        "Mara said the passage was perfectly safe, then the bridge failed."
    ),
)
record_test(
    "Deterministic reader story analysis links payoff, turns, loss, empathy, irony, subtext, and trust",
    any(item["kind"] == "payoff" for item in story_read["threads"])
    and any(item["kind"] == "yes_but" for item in story_read["turns"])
    and any(item["kind"] == "loss" for item in story_read["emotional_payoffs"])
    and any(item["kind"] == "empathy" for item in story_read["empathy_moments"])
    and any(item["kind"] == "subtext" for item in story_read["subtext_moments"])
    and any(item["kind"] == "irony" for item in story_read["irony_moments"])
    and any(item["kind"] == "trust_shift" for item in story_read["trust_shifts"])
    and all(
        item.get("setup_evidence") and item.get("payoff_evidence")
        for item in (
            [item for item in story_read["threads"] if item["kind"] != "setup"]
            + story_read["emotional_payoffs"]
            + story_read["empathy_moments"]
            + story_read["irony_moments"]
            + story_read["trust_shifts"]
        )
    ),
)
turn_kinds = {
    item["kind"] for item in _turns([
        "She found the key, but the gate failed.",
        "She lost the key, but found the old map.",
        "She found the key, and then opened the gate.",
        "She lost the key, and then the gate failed.",
    ])
}
record_test(
    "Story turn reader recognizes yes/no with but/and patterns",
    turn_kinds == {"yes_but", "no_but", "yes_and", "no_and"},
)
same_scene_loss = analyze_reader_story(
    [{
        "scene_id": "single-scene",
        "title": "The Last Walk",
        "content": "The child held the dog's paw all night. The dog died before dawn.",
    }],
    "single-scene",
    "The child held the dog's paw all night. The dog died before dawn.",
)
record_test(
    "Reader story analysis makes a direct, evidence-linked grief prediction",
    any(item["kind"] == "loss" for item in same_scene_loss["emotional_payoffs"])
    and "will land as grief" in same_scene_loss["emotional_payoffs"][0]["reaction"]
    and "held the dog's paw" in same_scene_loss["emotional_payoffs"][0]["setup_evidence"]
    and "dog died" in same_scene_loss["emotional_payoffs"][0]["payoff_evidence"],
)

noisy_text = (
    "I was standing dead center in the crossroads. He stopped dead and the mirror was held together by wire. "
    "Everyone loved to give that shadow credit. Frantically, my trembling hands shook. "
    "\"Did you hear that?\" a guard shrieked. Who gave the order to butcher my father in that alley?"
)
noisy_read = analyze_reader_story(
    [{"scene_id": "noisy", "title": "Chase", "content": noisy_text}], "noisy", noisy_text
)
record_test(
    "Reader story ignores idioms, physical fear, and guard dialogue instead of inventing grief or care",
    not noisy_read["emotional_payoffs"]
    and not noisy_read["empathy_moments"]
    and all("Did you hear" not in item["payoff_evidence"] for item in noisy_read["threads"])
    and any("father" in item["payoff_evidence"] for item in noisy_read["threads"]),
)

from backend.reader_feel import analyze_reader_feel

def _feel_scene():
    long_run = " ".join(["The corridor ran on past more gilded doors and statues, and the guards kept coming."] * 5)
    paragraphs = [
        "Don't get caught. Of all my rules for theft, it was the only one I hadn't broken until now.",
        "I walked like a cat on a hot roof, and I hated every minute of it. Why did tonight have to be the night?",
        long_run, "Run.", "Move or die.", "Hide.",
        "I had known it was tonight; in fact I'd bet my entire escape on it. " + long_run,
        "In theory.",
        "\"Perring is half-blind,\" the Captain hissed, trembling. \"Quiet! You want to summon the Knight?\"",
        "The guards were terrified. \"The Knight comes for us,\" the boy whispered, shaking.",
        "Lord Perring told them the Knight had cornered him. Of course, I thought; I was the one who did it.",
        "Our fear of the Knight grew. Then the Knight left, and the Knight was gone.",
        "Soon, the Knight was going to have to be afraid of me.",
    ]
    return "".join(f"<p>{p}</p>" for p in paragraphs) + "<p>" + "I kept going like a thief with a plan. " * 20 + "</p>"

feel_kinds = {item["kind"] for item in analyze_reader_feel(_feel_scene())}
record_test(
    "Reader feel estimates hook, undercut, fear of the unseen, and closing target from prose alone",
    {"hook", "undercut", "fear_of_unseen"}.issubset(feel_kinds),
    detail=str(feel_kinds),
)
record_test(
    "Reader feel stays silent on chapter headings and illustration tags",
    not analyze_reader_feel("<p>CHAPTER I.</p><p>[Illustration: a house]</p>" * 6),
)

import zipfile as _zipfile
from backend.routers.importer import _parse_import as _route_parse_import, _clean_user_path

with tempfile.TemporaryDirectory(prefix="settings_") as scriv_base_dir:
    scriv_base = Path(scriv_base_dir)
    modern = scriv_base / "Modern Novel.scriv"
    (modern / "Files" / "Data" / "B").mkdir(parents=True)
    (modern / "Novel.scrivx").write_text(
        '<ScrivenerProject><Binder><BinderItem UUID="A" Type="DraftFolder"><Title>Manuscript</Title><Children>'
        '<BinderItem UUID="B" Type="Text"><Title>Opening</Title></BinderItem></Children></BinderItem></Binder></ScrivenerProject>'
    )
    (modern / "Files" / "Data" / "B" / "content.rtf").write_text(r"{\rtf1\ansi Modern scene text here.\par}")
    legacy = scriv_base / "Legacy Novel.scriv"
    (legacy / "Files" / "Docs").mkdir(parents=True)
    (legacy / "Novel.scrivx").write_text(
        '<ScrivenerProject><Binder><BinderItem ID="1" Type="DraftFolder"><Title>Manuscript</Title><Children>'
        '<BinderItem ID="2" Type="Text"><Title>Opening</Title></BinderItem></Children></BinderItem></Binder></ScrivenerProject>'
    )
    (legacy / "Files" / "Docs" / "2.rtf").write_text(r"{\rtf1\ansi Legacy scene text here.\par}")
    mac_zip = scriv_base / "export.zip"
    with _zipfile.ZipFile(mac_zip, "w") as archive:
        archive.writestr("__MACOSX/Modern Novel.scriv/._Novel.scrivx", "stub")
        for item in modern.rglob("*"):
            archive.write(item, os.path.relpath(item, scriv_base))
    results = {}
    for label, target, uploaded in (
        ("folder", modern, False), ("scrivx", modern / "Novel.scrivx", False),
        ("legacy", legacy, False), ("zip", mac_zip, True),
    ):
        try:
            results[label] = _route_parse_import(str(target), str(target), uploaded)[2]["words"]
        except Exception as error:
            results[label] = str(error)
    try:
        _route_parse_import("/tmp/book.pdf", "book.pdf", True)
        pdf_rejected = False
    except ValueError:
        pdf_rejected = True
record_test(
    "Importer reads Scrivener folders, .scrivx paths, legacy projects, and macOS ZIPs, and rejects unsupported files",
    results == {"folder": 4, "scrivx": 4, "legacy": 4, "zip": 4} and pdf_rejected
    and _clean_user_path("'/Users/x/My\\ Novel.scriv' ") == "/Users/x/My Novel.scriv",
    detail=str(results),
)

import backend.wiktionary_lexicon as wiktionary
from backend.wiktionary_lexicon import expand_lexical_terms, get_lexicon_status, is_english_word

with tempfile.TemporaryDirectory() as wiktionary_temp:
    wiktionary_path = Path(wiktionary_temp) / "wiktionary.sqlite"
    source_path = Path(wiktionary_temp) / "source.sqlite"
    source_conn = sqlite3.connect(source_path)
    source_conn.executescript("""
        CREATE TABLE entries (id INTEGER PRIMARY KEY, word TEXT, lang TEXT);
        CREATE TABLE senses (id INTEGER PRIMARY KEY, entry_id INTEGER);
        CREATE TABLE glosses (sense_id INTEGER, gloss_text TEXT);
        CREATE TABLE synonyms (sense_id INTEGER, synonym_text TEXT);
        CREATE TABLE word_forms (entry_id INTEGER, form TEXT);
        INSERT INTO entries VALUES (1, 'love', 'English'), (2, 'cherish', 'English'), (3, 'cherished', 'English');
        INSERT INTO senses VALUES (10, 1), (20, 2);
        INSERT INTO glosses VALUES (10, 'to feel deep affection'), (20, 'to hold dear');
        INSERT INTO synonyms VALUES (10, 'cherish');
        INSERT INTO word_forms VALUES (2, 'cherished');
    """)
    source_conn.commit()
    source_conn.close()
    with patch.object(wiktionary, "WIKTIONARY_DB_PATH", wiktionary_path):
        wiktionary_install = wiktionary.install_wiktionary_database(source_path)
    love_terms = expand_lexical_terms({"love"}, wiktionary_path)
    lexical_story = analyze_reader_story(
        [{
            "scene_id": "lexical-scene",
            "title": "The Last Walk",
            "content": "The child cherished the dog. The dog died before dawn.",
        }],
        "lexical-scene",
        "The child cherished the dog. The dog died before dawn.",
        lexicon_path=wiktionary_path,
    )
    lexical_status = get_lexicon_status(wiktionary_path)
    english_word_found = is_english_word("love", wiktionary_path)
record_test(
    "Wiktionary SQLite adapter links synonyms and inflected forms into the reader lens",
    wiktionary_install["available"]
    and {"cherish", "cherished"}.issubset(love_terms)
    and english_word_found
    and any(item["kind"] == "loss" for item in lexical_story["emotional_payoffs"])
    and lexical_story["lexicon"]["available"]
    and lexical_status["license"] == "CC BY-SA 4.0"
    and "huggingface.co/datasets/cstr/en-wiktionary-sqlite-full" in lexical_status["source"],
    detail=(
        f"install={wiktionary_install['available']}, terms={love_terms}, "
        f"word={english_word_found}, "
        f"analysis_lexicon={lexical_story['lexicon']}, "
        f"payoffs={lexical_story['emotional_payoffs']}, status={lexical_status}"
    ),
)

from backend.main import app
from fastapi.testclient import TestClient

client = TestClient(app)

with tempfile.TemporaryDirectory() as candidate_api_temp:
    candidate_db_path = Path(candidate_api_temp) / "candidate_test.sqlite"
    init_db(candidate_db_path)
    candidate_conn = get_db(candidate_db_path)
    candidate_conn.executemany("""
        INSERT INTO binder_nodes (id, parent_id, node_type, title, content, sort_order)
        VALUES (?, ?, ?, ?, ?, ?)
    """, [
        ("candidate_book", None, "BOOK", "Candidate Book", "", 1),
        ("candidate_chapter", "candidate_book", "CHAPTER", "Chapter", "", 1),
        ("candidate_scene_a", "candidate_chapter", "SCENE", "Scene A", "Rowan met Aetherite beneath the Silver Archive.", 1),
        ("candidate_scene_b", "candidate_chapter", "SCENE", "Scene B", "They found Aetherite glowing beneath the Silver Archive.", 2),
    ])
    candidate_conn.commit()
    candidate_conn.close()
    with patch("backend.routers.entities.get_db", lambda: get_db(candidate_db_path)), \
         patch("backend.routers.entities.get_lexicon_status", return_value={"available": False}):
        candidate_response = client.get("/api/entities/candidates/candidate_book")
        register_alias_response = client.post(
            "/api/entities/quick-register",
            json={
                "name": "Mara Vale",
                "entity_type": "CHARACTER",
                "aliases": ["Mare", "Captain Vale", "The Fox"],
            },
        )
        alias_update_response = client.put(
            "/api/entities/mara_vale/aliases",
            json={"aliases": ["Mare", "Captain Vale", "The Fox", "Lady Vale"]},
        )
        conflicting_alias_response = client.post(
            "/api/entities/quick-register",
            json={
                "name": "Rowan Hale",
                "entity_type": "CHARACTER",
                "aliases": ["The Fox"],
            },
        )
candidate_api_result = candidate_response.json()
record_test(
    "Entity candidate endpoint scans book prose without auto-registering names",
    candidate_response.status_code == 200
    and {candidate["name"] for candidate in candidate_api_result["candidates"]}
    == {"Aetherite", "Silver Archive"}
    and not candidate_api_result["wiktionary_available"],
)
record_test(
    "Characters keep multiple aliases and reject conflicting identity names",
    register_alias_response.status_code == 200
    and alias_update_response.status_code == 200
    and alias_update_response.json()["aliases"]
    == ["Mare", "Captain Vale", "The Fox", "Lady Vale"]
    and conflicting_alias_response.status_code == 409,
)

with tempfile.TemporaryDirectory() as auto_temp:
    auto_db_path = Path(auto_temp) / "auto.sqlite"
    init_db(auto_db_path)
    auto_conn = get_db(auto_db_path)
    auto_conn.executemany("""
        INSERT INTO binder_nodes (id, parent_id, node_type, title, content, sort_order)
        VALUES (?, ?, ?, ?, ?, ?)
    """, [
        ("auto_book", None, "BOOK", "Auto Book", "", 1),
        ("auto_ch", "auto_book", "CHAPTER", "Chapter", "", 1),
        ("auto_a", "auto_ch", "SCENE", "A", "Rowan Hale walked to Kestrel Bay. Rowan Hale waited.", 1),
        ("auto_b", "auto_ch", "SCENE", "B", "Kestrel Bay was quiet. Rowan Hale left Kestrel Bay.", 2),
    ])
    auto_conn.commit()
    auto_conn.close()
    with patch("backend.routers.entities.get_db", lambda: get_db(auto_db_path)), \
         patch("backend.routers.entities.get_lexicon_status", return_value={"available": False}):
        auto_first = client.post("/api/entities/auto-detect/auto_book").json()
        auto_types = {item["name"]: item["entity_type"] for item in auto_first["added"]}
        auto_again = client.post("/api/entities/auto-detect/auto_book").json()
        auto_removed = client.delete("/api/entities/rowan_hale")
        auto_after_dismiss = client.post("/api/entities/auto-detect/auto_book").json()
        auto_confirm = client.post("/api/entities/kestrel_bay/confirm")
record_test(
    "Auto-detect tracks confident names, guesses type, never re-adds, and respects dismissals",
    auto_types == {"Rowan Hale": "CHARACTER", "Kestrel Bay": "LOCATION"}
    and auto_again["added"] == [] and auto_removed.status_code == 200
    and auto_after_dismiss["added"] == [] and auto_confirm.status_code == 200,
)

from backend.importing.text_cleaning import strip_rtf_native as _strip_rtf
_mac_rtf = _strip_rtf("{\\rtf1\\ansi\\cocoartf2639 \\f0 First paragraph.\\\nSecond paragraph.\\\nThird.}")
record_test(
    "RTF backslash-newline (Scrivener on Mac) becomes a paragraph break",
    [line for line in _mac_rtf.splitlines() if line.strip()] == ["First paragraph.", "Second paragraph.", "Third."],
)

with tempfile.TemporaryDirectory() as empty_scriv_temp:
    _empty_scriv = Path(empty_scriv_temp) / "Empty.scriv"
    _empty_scriv.mkdir()
    (_empty_scriv / "Empty.scrivx").write_text(
        '<?xml version="1.0"?><ScrivenerProject><Binder><BinderItem UUID="D" Type="DraftFolder"><Title>Manuscript</Title></BinderItem></Binder></ScrivenerProject>',
        encoding="utf-8",
    )
    _empty_res = client.post("/api/import/preview/local-path", json={"path": str(_empty_scriv)})
    (_empty_scriv / "Files" / "Data" / "X").mkdir(parents=True)
    (_empty_scriv / "Files" / "Data" / "X" / ".content.rtf.icloud").write_text("", encoding="utf-8")
    _icloud_res = client.post("/api/import/preview/local-path", json={"path": str(_empty_scriv)})
record_test(
    "Scrivener project with no text explains why, including iCloud placeholders",
    _empty_res.status_code == 400 and "No manuscript text" in _empty_res.json()["detail"]
    and _icloud_res.status_code == 400 and "iCloud" in _icloud_res.json()["detail"],
)

import time as _time
from backend.reader_story import analyze_reader_story as _analyze_story
_long_html = "".join("<p>Rowan asked why. She promised to return. He wondered if they would trust her.</p>" for _ in range(1500))
_t0 = _time.time()
_analyze_story([{"scene_id": "long", "title": "Long", "content": _long_html}], "long", _long_html)
record_test("Reader analysis stays fast on a very long single scene", _time.time() - _t0 < 15)

import sys as _sys, types as _types
with tempfile.TemporaryDirectory() as title_temp:
    title_db = Path(title_temp) / "t.sqlite"
    init_db(title_db)
    title_conn = get_db(title_db)
    _long = "<p>" + "The storm broke over the harbor and everyone ran. " * 3 + "</p>"
    title_conn.executemany("""
        INSERT INTO binder_nodes (id, parent_id, node_type, title, content, sort_order) VALUES (?, ?, ?, ?, ?, ?)
    """, [
        ("tb", None, "BOOK", "T", "", 1), ("tc", "tb", "CHAPTER", "C", "", 1),
        ("t1", "tc", "SCENE", "Scene 1", _long, 1), ("t2", "tc", "SCENE", "Opening", _long, 2),
        ("t3", "tc", "SCENE", "Untitled", _long, 3), ("t4", "tc", "SCENE", "Scene 4", "<p>Hi</p>", 4),
    ])
    title_conn.commit()
    title_conn.close()
    _fake = _types.ModuleType("ollama")
    _fake.chat = lambda **kwargs: {"message": {"content": "Storm Over Harbor"}}
    with patch.dict(_sys.modules, {"ollama": _fake}), patch("backend.routers.binder_tools.ENABLE_AI", True), \
         patch("backend.routers.binder_tools.get_db", lambda: get_db(title_db)):
        title_res = client.post("/api/binder/auto-title-scenes", json={"book_id": "tb"}).json()
record_test(
    "Auto-title retitles placeholder scenes and reports what it skipped",
    title_res["updated_count"] == 2 and title_res["skipped_named"] == 1 and title_res["skipped_short"] == 1,
)

health_res = client.get("/api/health")
record_test(
    "GET /api/health Reports Active AI Mode",
    health_res.status_code == 200
    and health_res.json().get("app") == "Thedral Studio"
    and health_res.json().get("mode") in {"LITE", "FULL"}
    and health_res.json().get("ai_enabled") == (health_res.json().get("mode") == "FULL"),
)

net_res = client.get("/api/network_info")
record_test("GET /api/network_info Returns Local LAN IP", net_res.status_code == 200 and "url" in net_res.json())

tree_res = client.get("/api/tree")
record_test("GET /api/tree Returns Hierarchical Binder", tree_res.status_code == 200 and isinstance(tree_res.json(), list))

entities_res = client.get("/api/entities")
record_test("GET /api/entities Returns Canonical Cast", entities_res.status_code == 200 and isinstance(entities_res.json(), list))

lore_res = client.get("/api/lore")
record_test("GET /api/lore Returns Living Lexicon", lore_res.status_code == 200 and isinstance(lore_res.json(), list))

api_books = tree_res.json()
if api_books:
    active_test_book = api_books[0]["id"]
    review_res = client.get(f"/api/publishing/review/{active_test_book}")
    record_test("GET /api/publishing/review Returns Manuscript Checks", review_res.status_code == 200 and "empty_scenes" in review_res.json())

    profile_res = client.get(f"/api/publishing/profile/{active_test_book}")
    record_test("GET /api/publishing/profile Returns Saved Export Settings", profile_res.status_code == 200 and "included_chapter_ids" in profile_res.json())

    proposals_res = client.get(f"/api/narrative/proposals?book_id={active_test_book}")
    record_test("GET /api/narrative/proposals Returns Pending Survey Suggestions", proposals_res.status_code == 200 and isinstance(proposals_res.json(), list))

    from backend.database import get_db
    feedback_conn = get_db()
    feedback_scene = feedback_conn.execute("""
        SELECT s.id
        FROM binder_nodes s JOIN binder_nodes ch ON ch.id = s.parent_id
        WHERE ch.parent_id = ? AND s.node_type = 'SCENE'
        ORDER BY ch.sort_order, s.sort_order LIMIT 1
    """, (active_test_book,)).fetchone()
    feedback_conn.close()
    if feedback_scene:
        story_read_res = client.post(
            f"/api/analytics/reader-story/{active_test_book}",
            json={
                "scene_id": feedback_scene["id"],
                "content": "Mara found the key, but the floodgate failed.",
            },
        )
        record_test(
            "Reader story analytics endpoint returns rule-based scene turns",
            story_read_res.status_code == 200
            and story_read_res.json().get("method", "").startswith("Transparent text-pattern rules")
            and any(turn.get("kind") == "yes_but" for turn in story_read_res.json().get("turns", [])),
        )
        story_chapter_res = client.post(
            f"/api/analytics/reader-story/{active_test_book}",
            json={"scene_id": "not_a_scene", "content": "Text."},
        )
        record_test(
            "Reader story for a non-scene node returns an empty result instead of an error",
            story_chapter_res.status_code == 200 and "not in the manuscript" in story_chapter_res.json().get("summary", ""),
        )
        imported_story_read = client.get(
            f"/api/analytics/reader-story/{active_test_book}/all"
        )
        imported_story_data = imported_story_read.json() if imported_story_read.status_code == 200 else {}
        imported_scene = next(
            (
                scene for scene in imported_story_data.get("analyses", [])
                if scene["scene_id"] == feedback_scene["id"]
            ),
            {},
        )
        record_test(
            "Whole-manuscript no-AI reader analysis returns cached scene results",
            imported_story_read.status_code == 200
            and imported_story_data.get("scene_count") == len(imported_story_data.get("analyses", []))
            and imported_scene.get("analysis", {}).get("scene_id") == feedback_scene["id"]
            and imported_scene.get("content") is not None
            and imported_scene.get("analysis", {}).get("method", "").startswith("Transparent text-pattern rules"),
        )

        prior_scene_id = f"reader_context_{uuid.uuid4().hex}"
        prior_chapter_id = f"reader_context_chapter_{uuid.uuid4().hex}"
        prior_scene_text = f"Earlier scene context marker {uuid.uuid4().hex}."
        context_conn = get_db()
        with context_conn:
            context_conn.execute("DELETE FROM binder_nodes WHERE id LIKE 'reader_context_%'")
            context_conn.execute("""
                INSERT INTO binder_nodes (id, parent_id, node_type, title, sort_order)
                VALUES (?, ?, 'CHAPTER', 'Reader context test chapter', -1000000000000)
            """, (prior_chapter_id, active_test_book))
            context_conn.execute("""
                INSERT INTO binder_nodes (id, parent_id, node_type, title, sort_order, content)
                VALUES (?, ?, 'SCENE', 'Reader context test scene', 0, ?)
            """, (prior_scene_id, prior_chapter_id, prior_scene_text))
        context_conn.close()
        live_scene_text = "<p>A live draft is compared with its neighboring scenes.</p>"
        analytics_res = client.post(
            f"/api/analytics/reader-experience/{active_test_book}",
            json={"scene_id": feedback_scene["id"], "content": live_scene_text},
        )
        analytics_data = analytics_res.json() if analytics_res.status_code == 200 else {}
        selected_analytics = next(
            (
                scene for scene in analytics_data.get("scenes", [])
                if scene["scene_id"] == feedback_scene["id"]
            ),
            {},
        )
        record_test(
            "Reader analytics compare live scene text in manuscript order",
            analytics_res.status_code == 200
            and analytics_data.get("selected_index", -1)
            == next(
                (
                    index for index, scene in enumerate(analytics_data.get("scenes", []))
                    if scene["scene_id"] == feedback_scene["id"]
                ),
                None,
            )
            and selected_analytics.get("telemetry", {}).get("words")
            == len(html_to_text(live_scene_text).split())
            and bool(selected_analytics.get("chapter_title")),
        )

        feedback_post = client.post("/api/reader-feedback", json={
            "scene_id": feedback_scene["id"],
            "revision_ref": "CURRENT",
            "reader_label": "Test reader",
            "engagement_rating": 4,
            "response": "The scene held my attention.",
        })
        feedback_id = feedback_post.json().get("id") if feedback_post.status_code == 200 else None
        scene_feedback = client.get(f"/api/reader-feedback/scene/{feedback_scene['id']}")
        book_feedback = client.get(f"/api/reader-feedback/book/{active_test_book}")
        feedback_visible = (
            feedback_id
            and scene_feedback.status_code == 200
            and any(entry["id"] == feedback_id for entry in scene_feedback.json())
            and book_feedback.status_code == 200
            and any(entry["id"] == feedback_id for entry in book_feedback.json())
        )
        if feedback_id:
            client.delete(f"/api/reader-feedback/{feedback_id}")
        record_test("Reader feedback is linked to its scene and manuscript", bool(feedback_visible))

        valid_ai_response = {
            "direct": [{
                "claim": "The reader sees that a passage is being validated.",
                "evidence": "sufficiently long passage",
            }],
            "reader_experience": {
                "emotional_arc": [{
                    "moment": "opening",
                    "estimate": "Cautious curiosity",
                    "level": 3,
                    "cause": "The phrasing withholds what is being validated.",
                    "evidence": "sufficiently long passage",
                }],
                "interest": [{
                    "moment": "opening",
                    "estimate": "Moderate interest",
                    "level": 3,
                    "cause": "The unexplained validation raises a question.",
                    "evidence": "validating the scanner response",
                }],
                "tension": [{
                    "moment": "middle",
                    "estimate": "Low tension",
                    "level": 2,
                    "cause": "The passage describes a calm action.",
                    "evidence": "The surrounding prose remains unchanged.",
                }],
                "mystery": [{
                    "moment": "opening",
                    "estimate": "Why the passage is being validated",
                    "level": 3,
                    "cause": "The purpose is not yet explained.",
                    "evidence": "validating the scanner response",
                }],
                "suspense": [],
                "character_connection": [],
                "confusion": [],
                "attention": [],
            },
        }
        sequential_ai_response = json.loads(json.dumps(valid_ai_response))
        sequential_ai_response["direct"][0]["evidence"] = prior_scene_text
        response_items = iter((
            [],
            {"direct": "not a list"},
            valid_ai_response,
            valid_ai_response,
            sequential_ai_response,
        ))
        ai_call_count = []
        mock_ollama = types.SimpleNamespace(
            chat=lambda **_kwargs: (
                ai_call_count.append(_kwargs)
                or {"message": {"content": json.dumps(next(response_items))}}
            )
        )
        with patch("backend.routers.reader_simulation.ENABLE_AI", True), patch.dict(sys.modules, {"ollama": mock_ollama}):
            invalid_root = client.post("/api/disclosures/scan", json={
                "scene_id": feedback_scene["id"],
                "text": "A sufficiently long passage for validating the scanner response.",
            })
            invalid_category = client.post("/api/disclosures/scan", json={
                "scene_id": feedback_scene["id"],
                "text": "A sufficiently long passage for validating the scanner response.",
            })
            previous_simulation_count = len(client.get(
                f"/api/reader-simulations/scene/{feedback_scene['id']}"
            ).json().get("simulations", []))
            simulation_payload = {
                "scene_id": feedback_scene["id"],
                "text": (
                    f"A sufficiently long passage for validating the scanner response {uuid.uuid4().hex}. "
                    + ("The surrounding prose remains unchanged. " * 120)
                ),
                "intended_takeaway": "The reader sees that a passage is being validated.",
            }
            first_simulation = client.post("/api/disclosures/scan", json=simulation_payload)
            repeated_simulation = client.post("/api/disclosures/scan", json=simulation_payload)
            revised_payload = {
                **simulation_payload,
                "text": simulation_payload["text"].replace(
                    "scanner response", "revised scanner response"
                ),
            }
            revised_simulation = client.post("/api/disclosures/scan", json=revised_payload)
            outside_excerpt_edit = client.post("/api/disclosures/scan", json={
                **revised_payload,
                "text": revised_payload["text"] + " An ending beat was added outside the analyzed excerpt.",
            })
            sequential_simulation = client.post("/api/disclosures/scan", json={
                **simulation_payload,
                "reading_mode": "sequential",
            })
            simulation_history = client.get(
                f"/api/reader-simulations/scene/{feedback_scene['id']}"
            )
        record_test(
            "Reader Horizon rejects malformed AI response shapes",
            invalid_root.json().get("status") == "error"
            and invalid_category.json().get("status") == "error"
        )
        first_saved_sim = first_simulation.json().get("simulation", {})
        repeated_saved_sim = repeated_simulation.json().get("simulation", {})
        snapshot_conn = get_db()
        persisted_source = snapshot_conn.execute(
            "SELECT source_text FROM reader_simulations WHERE id = ?",
            (first_saved_sim.get("id"),),
        ).fetchone()
        snapshot_conn.close()
        snapshot_checks = {
            "status": first_simulation.status_code == 200,
            "grounded excerpt": first_saved_sim.get("items", [{}])[0].get("source_excerpt") == "sufficiently long passage",
            "analyzed input": first_saved_sim.get("analyzed_text") == html_to_text(simulation_payload["text"])[:3500],
            "prompt/model/version": bool(
                first_saved_sim.get("system_prompt")
                and first_saved_sim.get("model_name")
                and first_saved_sim.get("prompt_version")
            ),
            "four new snapshots": len(simulation_history.json().get("simulations", [])) == previous_simulation_count + 4,
            "source length": first_saved_sim.get("source_characters") == len(html_to_text(simulation_payload["text"])),
            "source persisted": bool(
                persisted_source and persisted_source["source_text"] == html_to_text(simulation_payload["text"])
            ),
            "new result": first_simulation.json().get("reused") is False,
            "fresh reader contract": (
                first_saved_sim.get("prompt_version") == "first-read-v5"
                and first_saved_sim.get("intended_takeaway") == ""
                and "author's intended reader takeaway" not in first_saved_sim.get("system_prompt", "").lower()
                and "The reader sees that a passage is being validated."
                not in ai_call_count[2]["messages"][1]["content"]
            ),
            "reader experience estimates are grounded and typed": (
                any(
                    item.get("disclosure_type") == "EMOTIONAL_ARC"
                    and item.get("level") == 3
                    and item.get("cause")
                    and item.get("source_excerpt") == "sufficiently long passage"
                    for item in first_saved_sim.get("items", [])
                )
                and any(
                    item.get("disclosure_type") == "MYSTERY"
                    and item.get("level") == 3
                    and item.get("source_excerpt") == "validating the scanner response"
                    for item in first_saved_sim.get("items", [])
                )
            ),
        }
        record_test(
            "Reader simulation snapshots persist exact text and AI output",
            all(snapshot_checks.values()),
            str({name: value for name, value in snapshot_checks.items() if not value}),
        )
        record_test(
            "Unchanged text reuses the fixed fresh-reader simulation",
            first_saved_sim.get("id") == repeated_saved_sim.get("id")
            and repeated_simulation.json().get("reused") is True
            and revised_simulation.json().get("simulation", {}).get("id") != first_saved_sim.get("id")
            and revised_simulation.json().get("reused") is False
            and outside_excerpt_edit.json().get("reused") is True
            and outside_excerpt_edit.json().get("simulation", {}).get("reused_from_id")
            == revised_simulation.json().get("simulation", {}).get("id")
            and len(ai_call_count) == 5
        )
        sequential_saved = sequential_simulation.json().get("simulation", {})
        record_test(
            "Sequential reader rehearsal uses only preceding manuscript text",
            sequential_simulation.status_code == 200
            and sequential_saved.get("reading_mode") == "sequential"
            and prior_scene_text in sequential_saved.get("reading_context", "")
            and prior_scene_text in ai_call_count[4]["messages"][1]["content"]
            and sequential_saved.get("context_characters")
            == len(sequential_saved.get("reading_context", ""))
            and any(
                item.get("disclosure_type") == "DIRECT"
                and item.get("source_excerpt") == prior_scene_text
                and item.get("evidence_source") == "prior_context"
                for item in sequential_saved.get("items", [])
            ),
            str({
                "status": sequential_simulation.status_code,
                "reading_mode": sequential_saved.get("reading_mode"),
                "context": sequential_saved.get("reading_context"),
                "prompt_has_context": prior_scene_text in ai_call_count[4]["messages"][1]["content"],
                "items": sequential_saved.get("items"),
            }),
        )
        context_cleanup = get_db()
        with context_cleanup:
            context_cleanup.execute("DELETE FROM binder_nodes WHERE id = ?", (prior_scene_id,))
            context_cleanup.execute("DELETE FROM binder_nodes WHERE id = ?", (prior_chapter_id,))
        context_cleanup.close()
        cleanup_conn = get_db()
        with cleanup_conn:
            for simulation_id in (
                first_saved_sim.get("id"),
                revised_simulation.json().get("simulation", {}).get("id"),
                outside_excerpt_edit.json().get("simulation", {}).get("id"),
            ):
                if simulation_id:
                    cleanup_conn.execute("DELETE FROM reader_simulations WHERE id = ?", (simulation_id,))
        cleanup_conn.close()

        from backend.routers.craft import _reader_loop_signal
        repeated_observation = [{
            "narrative_claim": "The reader understands the clear message in the scene.",
        }]
        loop_fixture = [
            {
                "source_content_hash": f"revision-{number}",
                "analyzed_content_hash": f"analysis-{number}",
                "items": repeated_observation,
            }
            for number in range(3)
        ]
        record_test(
            "Revision-loop signal requires three distinct versions with repeated observations",
            _reader_loop_signal(loop_fixture)["active"]
            and not _reader_loop_signal(loop_fixture[:2])["active"]
        )

with tempfile.NamedTemporaryFile(suffix=".txt", mode="w", encoding="utf-8", delete=False) as preview_file:
    preview_file.write("A preview manuscript paragraph.")
    preview_path = preview_file.name
try:
    tree_before_preview = client.get("/api/tree").json()
    local_preview = client.post("/api/import/preview/local-path", json={"path": preview_path})
    upload_preview = client.post(
        "/api/import/preview/upload",
        files={"file": ("Preview.txt", b"A preview manuscript paragraph.", "text/plain")},
    )
    tree_after_preview = client.get("/api/tree").json()
    record_test(
        "Import previews report content without changing the Binder",
        local_preview.status_code == 200
        and local_preview.json()["report"]["words"] == 4
        and upload_preview.status_code == 200
        and upload_preview.json()["report"]["scenes"] == 1
        and tree_before_preview == tree_after_preview
    )
finally:
    os.remove(preview_path)

# =============================================================================
# 9. TEXT-PATTERN SIGNAL ANALYSIS
# =============================================================================
print(f"\n{BOLD}9. Text-pattern Signal Analysis{RESET}")

from backend.spectrometer import analyze_text_signals

signal_fixture = (
    "The blade struck the iron gate."
    "\n\nHis gaze lingered as their fingers touched."
    "\n\nA hidden secret changed everything."
    "\n\nRain tapped stone beneath the cold night sky."
)
signals = analyze_text_signals(signal_fixture)
record_test("Counts selected action cues", signals["signals"]["action_cues_per_1000_words"]["count"] > 0)
record_test("Counts selected relational cues", signals["signals"]["relational_cues_per_1000_words"]["count"] > 0)
record_test("Counts selected mystery cues", signals["signals"]["mystery_cues_per_1000_words"]["count"] > 0)
record_test("Counts selected sensory cues", signals["signals"]["sensory_cues_per_1000_words"]["count"] > 0)
record_test("Signal report disclaims reader-response measurement", "do not measure reader emotions" in signals["limitations"])


print("\n10. AI helper integration")
import types as _types
from backend.ai_helper import spelling_findings, find_name_candidates, ai_meta, deterministic_brief
from backend.routers import lore as lore_router

record_test("Misspelled lore term is caught without AI",
            spelling_findings([{"term": "Aethermoor"}], "They rode toward Aethermor at dawn.")[0]["violation"] == "Aethermor")
record_test("Correct lore spelling is not flagged", spelling_findings([{"term": "Aethermoor"}], "Aethermoor stood.") == [])
record_test("Invented names are found by text scan",
            "Zhivrek" in find_name_candidates("The old guard feared Zhivrek and Zhivrek knew it.", set()))
record_test("ai_meta labels engine", ai_meta(True)["engine"] == "local-ai" and ai_meta(False)["engine"] == "deterministic")

_lore_text = "The council met at dusk. " * 6 + "Guards saw Zhivrek watching from the tower."
with patch("backend.routers.lore.ENABLE_AI", False):
    _lite = lore_router.auto_extract_lore(lore_router.ExtractLoreRequest(scene_text=_lore_text))
record_test("Lore extraction works in Lite mode and says no AI",
            _lite["ai_generated"] is False and any(p["term"] == "Zhivrek" for p in _lite["proposed_lore"]))

_prompts = []
def _fake_chat(**kwargs):
    _prompts.append(kwargs["messages"][0]["content"])
    return {"message": {"content": '[{"term": "Zhivrek", "category": "LORE", "rule_definition": "A watcher."}]'}}
with patch("backend.routers.lore.ENABLE_AI", True), patch.dict(sys.modules, {"ollama": _types.SimpleNamespace(chat=_fake_chat)}):
    _ai = lore_router.auto_extract_lore(lore_router.ExtractLoreRequest(scene_text=_lore_text))
record_test("AI extraction is labeled and receives text-scan candidates",
            _ai["ai_generated"] is True and "Zhivrek" in _prompts[0] and _ai["proposed_lore"][0]["source"] == "ai")

def _broken_chat(**kwargs):
    raise RuntimeError("down")
with patch("backend.routers.lore.ENABLE_AI", True), patch.dict(sys.modules, {"ollama": _types.SimpleNamespace(chat=_broken_chat)}):
    _fail = lore_router.auto_extract_lore(lore_router.ExtractLoreRequest(scene_text=_lore_text))
record_test("AI failure falls back to text scan, not fake findings",
            _fail["ai_generated"] is False and _fail["ai_error"] and _fail["proposed_lore"])

print("\n11. Feature inventory guard")
from backend.main import app as _guard_app
_live_routes = {f"{m.upper()} {p}" for p, ops in _guard_app.openapi()["paths"].items() for m in ops}
_expected_routes = [line for line in open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "route_manifest.txt")).read().splitlines() if line]
_missing_routes = [r for r in _expected_routes if r not in _live_routes]
record_test("Every recorded API route is still served", not _missing_routes, str(_missing_routes[:10]))

import importlib as _importlib
_missing_names = []
for _line in open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "api_manifest.txt")).read().splitlines():
    _mod, _, _name = _line.partition(":")
    if not hasattr(_importlib.import_module(_mod), _name):
        _missing_names.append(_line)
record_test("Every recorded module function is still importable from its original path", not _missing_names, str(_missing_names[:10]))

# =============================================================================
# SUMMARY REPORT
_promise_book = client.get("/api/tree").json()[0]["id"]
_promise_post = client.post("/api/narrative/promise", json={
    "book_id": _promise_book, "scene_id": "scene_1", "promise_desc": "Locked drawer", "target_tick": 7.0,
})
_promise_rows = client.get(f"/api/narrative/promises?book_id={_promise_book}").json()
_promise_row = next((p for p in _promise_rows if p["promise_desc"] == "Locked drawer"), {})
_promise_paid = client.put(f"/api/narrative/promises/{_promise_row.get('id', 0)}", json={"status": "PAID_OFF"})
_promise_bad = client.put(f"/api/narrative/promises/{_promise_row.get('id', 0)}", json={"status": "NOPE"})
_promise_after = next((p for p in client.get(f"/api/narrative/promises?book_id={_promise_book}").json() if p["id"] == _promise_row.get("id")), {})
_promise_del = client.delete(f"/api/narrative/promises/{_promise_row.get('id', 0)}")
_promise_gone = client.delete(f"/api/narrative/promises/{_promise_row.get('id', 0)}")
record_test(
    "Promise ledger: plant, list, resolve, reject bad status, delete",
    _promise_post.status_code == 200 and _promise_row.get("planted_scene_id") == "scene_1"
    and _promise_paid.status_code == 200 and _promise_bad.status_code == 400
    and _promise_after.get("status") == "PAID_OFF"
    and _promise_del.status_code == 200 and _promise_gone.status_code == 404,
)

_rel_a = client.post("/api/entities/quick-register", json={"name": "Rel Alpha", "entity_type": "CHARACTER"})
_rel_b = client.post("/api/entities/quick-register", json={"name": "Rel Beta", "entity_type": "CHARACTER"})
_rel_ents = {e["name"]: e["entity_id"] for e in client.get("/api/entities").json()}
_rel_ok = client.post("/api/narrative/connect", json={"source_id": _rel_ents["Rel Alpha"], "target_id": _rel_ents["Rel Beta"], "rel_type": "allied", "description": "Old friends"})
_rel_dup = client.post("/api/narrative/connect", json={"source_id": _rel_ents["Rel Alpha"], "target_id": _rel_ents["Rel Beta"], "rel_type": "ALLIED"})
_rel_self = client.post("/api/narrative/connect", json={"source_id": _rel_ents["Rel Alpha"], "target_id": _rel_ents["Rel Alpha"], "rel_type": "ALLIED"})
_rel_missing = client.post("/api/narrative/connect", json={"source_id": "nope", "target_id": _rel_ents["Rel Beta"], "rel_type": "ALLIED"})
_rel_list = [r for r in client.get("/api/narrative/connections").json() if r["source_name"] == "Rel Alpha"]
_rel_del = client.delete(f"/api/narrative/connections/{_rel_ok.json().get('rel_id')}")
_rel_del2 = client.delete(f"/api/narrative/connections/{_rel_ok.json().get('rel_id')}")
record_test(
    "Relationships: create, dedupe, validate, list with names, delete",
    _rel_ok.status_code == 200 and _rel_dup.json().get("rel_id") == _rel_ok.json().get("rel_id")
    and _rel_self.status_code == 400 and _rel_missing.status_code == 404
    and len(_rel_list) == 1 and _rel_list[0]["rel_type"] == "ALLIED" and _rel_list[0]["target_name"] == "Rel Beta"
    and _rel_del.status_code == 200 and _rel_del2.status_code == 404,
)
for _n in ("Rel Alpha", "Rel Beta"):
    client.delete(f"/api/entities/{_rel_ents[_n]}")

_ser_book = client.post("/api/node", json={"node_type": "BOOK", "title": "Serendipity Book"}).json()
_ser_ch = client.post("/api/node", json={"node_type": "CHAPTER", "title": "C", "parent_id": _ser_book["id"]}).json()
_ser_s1 = client.post("/api/node", json={"node_type": "SCENE", "title": "S1", "parent_id": _ser_ch["id"]}).json()
_ser_s2 = client.post("/api/node", json={"node_type": "SCENE", "title": "S2", "parent_id": _ser_ch["id"]}).json()
client.post("/api/lore", json={"term": "Glimmerstone", "category": "ITEM", "rule_definition": "Glows near lies"})
client.put(f"/api/node/{_ser_s1['id']}", json={"content": "<p>The Glimmerstone flared on the table.</p>"})
client.put(f"/api/node/{_ser_s2['id']}", json={"content": "<p>Rain fell on an empty street.</p>"})
with patch("backend.routers.narrative.ENABLE_AI", False):
    _ser_dormant = client.get(f"/api/narrative/serendipity/{_ser_s2['id']}").json()
    _ser_present = client.get(f"/api/narrative/serendipity/{_ser_s1['id']}").json()
    _ser_404 = client.get("/api/narrative/serendipity/missing")
record_test(
    "Serendipity: finds present and dormant lore without AI",
    any(i["label"] == "Glimmerstone" for i in _ser_dormant["dormant"])
    and any(i["label"] == "Glimmerstone" for i in _ser_present["present"])
    and _ser_dormant["ai_generated"] is False and _ser_404.status_code == 404,
)
client.delete(f"/api/node/{_ser_book['id']}")
client.delete("/api/lore/Glimmerstone")

import sqlite3 as _sq, tempfile as _tf
from backend.wiktionary_lexicon import expand_lexical_terms as _expand, get_lexicon_status as _lex_status
_lex_dir = _tf.mkdtemp()
_lex_path = os.path.join(_lex_dir, "wikt.sqlite")
_lx = _sq.connect(_lex_path)
_lx.executescript("""
CREATE TABLE entries (id INTEGER PRIMARY KEY, word TEXT, title TEXT, redirect TEXT, pos TEXT, pos_title TEXT, lang_code TEXT, lang TEXT, etymology_text TEXT);
CREATE TABLE senses (id INTEGER PRIMARY KEY, entry_id INTEGER, sense_index INTEGER);
CREATE TABLE glosses (id INTEGER PRIMARY KEY, sense_id INTEGER, gloss_text TEXT, gloss_order INTEGER);
CREATE TABLE forms (id INTEGER PRIMARY KEY, entry_id INTEGER, form_text TEXT, sense_index INTEGER);
CREATE TABLE synonyms (id INTEGER PRIMARY KEY, entry_id INTEGER, synonym_word TEXT, sense_index INTEGER);
INSERT INTO entries (id, word, pos, lang_code, lang) VALUES (1, 'Kill', 'verb', 'en', 'English');
INSERT INTO forms (entry_id, form_text) VALUES (1, 'killed'), (1, 'Kills');
INSERT INTO synonyms (entry_id, synonym_word) VALUES (1, 'slay'), (1, 'I-495');
""")
_lx.commit(); _lx.close()
_lex_terms = _expand(["kill", "whisper"], _lex_path)
record_test(
    "Wiktionary: real-schema synonyms and forms expand; gaps get regular forms; noise dropped",
    _lex_status(_lex_path)["available"]
    and {"slay", "killed", "whispered", "whispering"} <= _lex_terms
    and "i-495" not in _lex_terms and "kill" not in _lex_terms,
)

_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_fe_missing = []
for _line in open(os.path.join(_root, "tests", "frontend_manifest.txt")).read().splitlines():
    _kind, _, _val = _line.partition(":")
    if _kind == "file" and not os.path.exists(os.path.join(_root, _val)):
        _fe_missing.append(_line)
    elif _kind == "action" and f" {_val}(" not in open(os.path.join(_root, "frontend/src/storeActions.js")).read():
        _fe_missing.append(_line)
record_test("Every recorded frontend store action and component file still exists", not _fe_missing, str(_fe_missing[:10]))

# =============================================================================
print(f"\n{BOLD}{HEADER}==============================================================================={RESET}")
if tests_passed == tests_run:
    print(f"{BOLD}\033[92m✨ MASTER VERIFICATION COMPLETE: ALL {tests_run}/{tests_run} TESTS PASSED!{RESET}")
    print(f"{BOLD}Thedral sovereign core is stable, verified, and ready for Git.{RESET}")
else:
    print(f"{BOLD}\033[91m⚠️ VERIFICATION INCOMPLETE: {tests_passed}/{tests_run} passed.{RESET}")
print(f"{BOLD}{HEADER}==============================================================================={RESET}\n")